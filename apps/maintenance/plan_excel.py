"""Planes preventivos por Excel (PM-08): exportar la versión vigente e importar con vista previa.

El archivo tiene tres hojas: «Planes» (encabezado y frecuencia), «Checklist»
y «Recursos», unidas por el código del plan. Lo que se exporta se puede
volver a importar tal cual.

Importar nunca modifica una versión publicada (ADR-016): un código nuevo crea
el plan con su versión 1 publicada; un código existente con contenido
distinto crea una versión borrador para revisar y publicar; sin cambios, se
omite. El nombre y la clase de un plan existente no cambian desde el archivo.
"""

import datetime
import io
import zipfile
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.translation import gettext, gettext_lazy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from apps.accounts.permissions import Perm, require_perm
from apps.assets.imports import MAX_FILE_BYTES, ImportFileError, _cell_text, _normalize
from apps.assets.models import METER_DEFAULT_UNITS, AssetClass, MeterKind
from apps.audit import services as audit
from apps.audit.models import AuditEvent

from .models import CalendarUnit, ItemType, MaintenancePlan, PlanResource, PlanTaskItem, PlanVersion
from .services import bulk_create_published_plans, create_draft_with_content

MAX_PLANS = 500
MAX_ITEMS_PER_PLAN = 100

PLANS_SHEET = gettext_lazy("Planes")
ITEMS_SHEET = gettext_lazy("Checklist")
RESOURCES_SHEET = gettext_lazy("Recursos")

# (clave, encabezado, obligatoria, ayuda)
PLAN_COLUMNS = [
    (
        "code",
        gettext_lazy("Código"),
        True,
        gettext_lazy("Único por plan. Si ya existe, se crea un borrador."),
    ),
    ("name", gettext_lazy("Nombre"), True, ""),
    ("asset_class", gettext_lazy("Clase"), False, gettext_lazy("Código o nombre de una clase existente.")),
    ("calendar_interval", gettext_lazy("Cada (calendario)"), False, gettext_lazy("Número entero, ej: 3.")),
    ("calendar_unit", gettext_lazy("Unidad de calendario"), False, gettext_lazy("días, semanas o meses.")),
    ("meter_kind", gettext_lazy("Medidor"), False, gettext_lazy("Horómetro, Odómetro o Ciclos.")),
    ("meter_interval", gettext_lazy("Cada (medidor)"), False, gettext_lazy("Ej: 250 (horas) o 10000 (km).")),
    ("schedule_mode", gettext_lazy("Reprogramación"), False, gettext_lazy("Flotante (por defecto) o Fijo.")),
    ("tolerance_pct", gettext_lazy("Tolerancia (%)"), False, gettext_lazy("Por defecto 10.")),
    ("estimated_minutes", gettext_lazy("Duración estimada (min)"), False, gettext_lazy("Por defecto 60.")),
    ("instructions", gettext_lazy("Instrucciones"), False, ""),
    ("version", gettext_lazy("Versión vigente"), False, gettext_lazy("Informativa: se ignora al importar.")),
]
ITEM_COLUMNS = [
    ("plan", gettext_lazy("Código del plan"), True, ""),
    ("order", gettext_lazy("Orden"), False, gettext_lazy("Por defecto, el orden de las filas.")),
    ("text", gettext_lazy("Tarea"), True, ""),
    (
        "item_type",
        gettext_lazy("Tipo"),
        False,
        gettext_lazy("Verificación (por defecto), Medición, Texto, Foto o Firma."),
    ),
    ("unit", gettext_lazy("Unidad"), False, gettext_lazy("Solo mediciones, ej: bar.")),
    ("min_value", gettext_lazy("Mínimo"), False, gettext_lazy("Solo mediciones.")),
    ("max_value", gettext_lazy("Máximo"), False, gettext_lazy("Solo mediciones.")),
    ("required", gettext_lazy("Obligatorio"), False, gettext_lazy("Sí (por defecto) o No.")),
]
RESOURCE_COLUMNS = [
    ("plan", gettext_lazy("Código del plan"), True, ""),
    ("kind", gettext_lazy("Tipo"), False, gettext_lazy("Repuesto (por defecto) o Herramienta.")),
    ("description", gettext_lazy("Descripción"), True, ""),
    ("quantity", gettext_lazy("Cantidad"), False, gettext_lazy("Por defecto 1.")),
    ("unit", gettext_lazy("Unidad"), False, ""),
    (
        "part",
        gettext_lazy("Código de repuesto"),
        False,
        gettext_lazy("Opcional: código de un repuesto del catálogo de inventario."),
    ),
]
SHEETS = [
    ("plans", PLANS_SHEET, PLAN_COLUMNS),
    ("items", ITEMS_SHEET, ITEM_COLUMNS),
    ("resources", RESOURCES_SHEET, RESOURCE_COLUMNS),
]
YES = {"si", "s", "yes", "y", "true", "1", "x"}
NO = {"no", "n", "false", "0"}


def _aliases(choices, extra=None) -> dict:
    """Texto de celda normalizado → valor: acepta el valor interno, la etiqueta y su primera palabra."""
    index = {}
    for value, label in choices:
        for alias in (value, str(label), str(label).split(" ")[0].split("(")[0]):
            index.setdefault(_normalize(alias), value)
    index.update(extra or {})
    return index


def _choice_aliases():
    meter_extra = {_normalize(unit): kind for kind, unit in METER_DEFAULT_UNITS.items()}
    meter_extra.update({"horas": MeterKind.HOURS, "kilometraje": MeterKind.KM})
    return {
        "calendar_unit": _aliases(
            CalendarUnit.choices,
            {"dia": "days", "semana": "weeks", "mes": "months", "d": "days", "sem": "weeks"},
        ),
        "meter_kind": _aliases([c for c in MeterKind.choices if c[0] != MeterKind.OTHER], meter_extra),
        "schedule_mode": _aliases(PlanVersion.ScheduleMode.choices),
        "item_type": _aliases(ItemType.choices),
        "kind": _aliases(PlanResource.Kind.choices),
    }


# --- Exportar ----------------------------------------------------------------


def _number_cell(value):
    if value is None:
        return None
    value = Decimal(value)
    return int(value) if value == value.to_integral_value() else float(value)


def export_rows(*, membership):
    """Filas de las tres hojas con la versión vigente de cada plan (consultas acotadas)."""
    require_perm(membership, Perm.PLAN_MANAGE)
    plans = list(
        MaintenancePlan.objects.filter(current_version__isnull=False)
        .select_related("asset_class")
        .order_by("code")
    )
    version_ids = [p.current_version_id for p in plans]
    items, resources = {}, {}
    for item in PlanTaskItem.objects.filter(version_id__in=version_ids).order_by("order"):
        items.setdefault(item.version_id, []).append(item)
    for resource in (
        PlanResource.objects.filter(version_id__in=version_ids)
        .select_related("part")
        .order_by("kind", "description")
    ):
        resources.setdefault(resource.version_id, []).append(resource)
    plan_rows, item_rows, resource_rows = [], [], []
    for plan in plans:
        v = plan.current_version
        plan_rows.append(
            [
                plan.code,
                plan.name,
                plan.asset_class.code if plan.asset_class else "",
                v.calendar_interval,
                str(v.get_calendar_unit_display()) if v.calendar_interval else "",
                str(v.get_meter_kind_display()) if v.meter_kind else "",
                _number_cell(v.meter_interval),
                str(v.get_schedule_mode_display()),
                v.tolerance_pct,
                v.estimated_minutes,
                v.instructions,
                v.number,
            ]
        )
        for item in items.get(v.id, ()):
            item_rows.append(
                [
                    plan.code,
                    item.order,
                    item.text,
                    str(item.get_item_type_display()),
                    item.unit,
                    _number_cell(item.min_value),
                    _number_cell(item.max_value),
                    gettext("Sí") if item.required else gettext("No"),
                ]
            )
        for resource in resources.get(v.id, ()):
            resource_rows.append(
                [
                    plan.code,
                    str(resource.get_kind_display()),
                    resource.description,
                    _number_cell(resource.quantity),
                    resource.unit,
                    resource.part.code if resource.part_id else "",
                ]
            )
    return {"plans": plan_rows, "items": item_rows, "resources": resource_rows}


def build_plans_workbook(*, membership) -> bytes:
    rows = export_rows(membership=membership)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="0B2E4F")
    for key, title, columns in SHEETS:
        ws = wb.create_sheet(str(title))
        for col, (_key, label, _required, _help) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col, value=str(label))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            ws.column_dimensions[cell.column_letter].width = max(14, len(str(label)) + 6)
        for row in rows[key]:
            ws.append(row)
        ws.freeze_panes = "A2"
    help_sheet = wb.create_sheet(gettext("Instrucciones"))
    help_sheet.column_dimensions["A"].width = 26
    help_sheet.column_dimensions["B"].width = 90
    help_sheet.append([gettext("Columna"), gettext("Cómo llenarla")])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for _key, title, columns in SHEETS:
        help_sheet.append([str(title)])
        help_sheet.cell(row=help_sheet.max_row, column=1).font = Font(bold=True)
        for _k, label, required, help_text in columns:
            column = gettext("%(column)s (obligatoria)") % {"column": label} if required else str(label)
            help_sheet.append([column, str(help_text)])
    help_sheet.append([])
    for line in (
        gettext("Un código nuevo crea el plan con su versión 1 publicada."),
        gettext(
            "Un código existente con cambios crea una versión borrador para revisar y publicar; "
            "la versión publicada no se modifica."
        ),
        gettext("Antes de guardar verás una vista previa con los errores de cada fila."),
    ):
        help_sheet.append(["", line])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_filename() -> str:
    return f"planes-{datetime.date.today():%Y%m%d}.xlsx"


# --- Leer el archivo ---------------------------------------------------------


def _read_sheet(workbook, title, columns, *, required):
    name = str(title)
    match = next((s for s in workbook.sheetnames if _normalize(s) == _normalize(name)), None)
    if match is None:
        if required:
            raise ImportFileError(gettext("Falta la hoja «%(sheet)s».") % {"sheet": name})
        return []
    rows = workbook[match].iter_rows(values_only=True)
    header = next(rows, None) or ()
    aliases = {}
    for key, label, _required, _help in columns:
        aliases[_normalize(label)] = key
        aliases[_normalize(key)] = key
    keys = [aliases.get(_normalize(h)) if h is not None else None for h in header]
    missing = [str(label) for key, label, req, _ in columns if req and key not in keys]
    if missing:
        raise ImportFileError(
            gettext("Faltan columnas en la hoja «%(sheet)s»: %(columns)s.")
            % {"sheet": name, "columns": ", ".join(missing)}
        )
    parsed = []
    for number, values in enumerate(rows, start=2):
        row = {key: _cell_text(v) for key, v in zip(keys, values, strict=False) if key}
        if any(row.values()):
            row["_row"] = number
            parsed.append(row)
    return parsed


def parse_plan_workbook(uploaded_file) -> dict:
    """Filas de texto de las hojas Planes, Checklist y Recursos (esta última es opcional)."""
    if uploaded_file.size > MAX_FILE_BYTES:
        raise ImportFileError(gettext("El archivo supera 5 MB."))
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ImportFileError(gettext("No se pudo leer el archivo. Usa un Excel .xlsx.")) from exc
    try:
        data = {
            "plans": _read_sheet(workbook, PLANS_SHEET, PLAN_COLUMNS, required=True),
            "items": _read_sheet(workbook, ITEMS_SHEET, ITEM_COLUMNS, required=True),
            "resources": _read_sheet(workbook, RESOURCES_SHEET, RESOURCE_COLUMNS, required=False),
        }
    finally:
        workbook.close()
    if not data["plans"]:
        raise ImportFileError(gettext("La hoja «Planes» no tiene filas con datos."))
    if len(data["plans"]) > MAX_PLANS:
        raise ImportFileError(
            gettext("El archivo tiene más de %(max)s planes; divídelo en partes.") % {"max": MAX_PLANS}
        )
    return data


# --- Vista previa (sin escribir) ------------------------------------------------


def _decimal(text: str):
    """Número con coma o punto decimal (es-CL: 1.234,5)."""
    text = text.replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(text) from exc
    if not value.is_finite():
        raise ValueError(text)
    return value


def _integer(text: str):
    value = _decimal(text)
    if value != value.to_integral_value():
        raise ValueError(text)
    return int(value)


@dataclass
class PlanRow:
    row: int
    code: str
    name: str
    action: str  # create | draft | unchanged | error
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    data: dict = field(default_factory=dict)
    item_count: int = 0
    frequency_label: str = ""


@dataclass
class PlanPreview:
    rows: list[PlanRow]
    # Filas de Checklist o Recursos cuyo plan no está en la hoja Planes.
    orphans: list[str]

    @property
    def valid_rows(self):
        return [r for r in self.rows if r.action in ("create", "draft")]

    def count(self, action):
        return sum(1 for r in self.rows if r.action == action)

    def to_session(self) -> dict:
        return {"rows": [asdict(r) for r in self.valid_rows]}


def _parse_plan(raw, aliases, row: PlanRow) -> dict:
    """Campos de la versión desde la fila de «Planes» (errores en `row.errors`)."""
    fields = {
        "calendar_interval": None,
        "calendar_unit": CalendarUnit.DAYS,
        "meter_kind": "",
        "meter_interval": None,
    }

    def pick(key, default):
        text = raw.get(key, "")
        if not text:
            return default
        value = aliases[key].get(_normalize(text))
        if value is None:
            row.errors.append(gettext("«%(value)s» no es un valor válido.") % {"value": text})
        return value or default

    try:
        if raw.get("calendar_interval"):
            fields["calendar_interval"] = _integer(raw["calendar_interval"])
            if fields["calendar_interval"] < 1:
                raise ValueError
            fields["calendar_unit"] = pick("calendar_unit", CalendarUnit.DAYS)
    except ValueError:
        row.errors.append(gettext("«Cada (calendario)» debe ser un entero mayor que 0."))
    if raw.get("meter_kind") or raw.get("meter_interval"):
        fields["meter_kind"] = pick("meter_kind", "")
        try:
            fields["meter_interval"] = _decimal(raw.get("meter_interval", ""))
            if fields["meter_interval"] <= 0:
                raise ValueError
        except ValueError:
            row.errors.append(gettext("«Cada (medidor)» debe ser un número mayor que 0."))
        if not fields["meter_kind"]:
            row.errors.append(gettext("Para disparar por medidor indica el tipo de medidor."))
    if not fields["calendar_interval"] and not fields["meter_interval"]:
        row.errors.append(gettext("Define una frecuencia por calendario, por medidor o ambas."))
    fields["schedule_mode"] = pick("schedule_mode", PlanVersion.ScheduleMode.FLOATING)
    for key, default, low, high, label in (
        ("tolerance_pct", 10, 0, 100, gettext("Tolerancia (%)")),
        ("estimated_minutes", 60, 1, 100000, gettext("Duración estimada (min)")),
    ):
        try:
            value = _integer(raw[key]) if raw.get(key) else default
            if not low <= value <= high:
                raise ValueError
            fields[key] = value
        except ValueError:
            row.errors.append(
                gettext("«%(column)s» debe ser un entero entre %(low)s y %(high)s.")
                % {"column": label, "low": low, "high": high}
            )
    fields["instructions"] = raw.get("instructions", "")
    return fields


def _parse_item(raw, aliases, errors) -> dict | None:
    where = gettext("Checklist fila %(row)s") % {"row": raw["_row"]}
    text = raw.get("text", "")
    if len(text) > 250:
        errors.append(gettext("%(where)s: la tarea supera 250 caracteres.") % {"where": where})
        return None
    item_type = aliases["item_type"].get(_normalize(raw.get("item_type") or "check"))
    if item_type is None:
        errors.append(
            gettext("%(where)s: tipo «%(value)s» no válido.") % {"where": where, "value": raw["item_type"]}
        )
        return None
    item = {"text": text, "item_type": item_type, "unit": "", "min_value": None, "max_value": None}
    required = _normalize(raw.get("required", "")) or "si"
    if required not in YES | NO:
        errors.append(gettext("%(where)s: «Obligatorio» debe ser Sí o No.") % {"where": where})
        return None
    item["required"] = required in YES
    if item_type == ItemType.MEASUREMENT:
        item["unit"] = raw.get("unit", "")[:20]
        try:
            for key in ("min_value", "max_value"):
                item[key] = _decimal(raw[key]) if raw.get(key) else None
        except ValueError:
            errors.append(gettext("%(where)s: mínimo y máximo deben ser números.") % {"where": where})
            return None
        if (
            item["min_value"] is not None
            and item["max_value"] is not None
            and item["min_value"] > item["max_value"]
        ):
            errors.append(
                gettext("%(where)s: el máximo no puede ser menor que el mínimo.") % {"where": where}
            )
            return None
    try:
        order = _integer(raw["order"]) if raw.get("order") else raw["_row"]
    except ValueError:
        order = raw["_row"]
    return {"_order": order, **item}


def _parse_resource(raw, aliases, errors, parts=None) -> dict | None:
    where = gettext("Recursos fila %(row)s") % {"row": raw["_row"]}
    kind = aliases["kind"].get(_normalize(raw.get("kind") or "part"))
    if kind is None:
        errors.append(
            gettext("%(where)s: tipo «%(value)s» no válido.") % {"where": where, "value": raw["kind"]}
        )
        return None
    try:
        quantity = _decimal(raw["quantity"]) if raw.get("quantity") else Decimal(1)
        if quantity <= 0:
            raise ValueError
    except ValueError:
        errors.append(gettext("%(where)s: la cantidad debe ser un número mayor que 0.") % {"where": where})
        return None
    part_id = None
    if raw.get("part"):  # repuesto del catálogo (INV-11), opcional
        part = (parts or {}).get(raw["part"].strip().lower())
        if part is None:
            errors.append(
                gettext("%(where)s: el repuesto «%(value)s» no existe en el catálogo.")
                % {"where": where, "value": raw["part"]}
            )
            return None
        if kind != PlanResource.Kind.PART:
            errors.append(gettext("%(where)s: solo un repuesto se enlaza al catálogo.") % {"where": where})
            return None
        part_id = str(part.pk)
    return {
        "kind": kind,
        "description": raw["description"][:200],
        "quantity": quantity,
        "unit": raw.get("unit", "")[:20],
        "part_id": part_id,
    }


def _content_key(fields, items, resources):
    """Contenido comparable de una versión (para detectar filas sin cambios)."""
    normalized = tuple(
        (k, Decimal(fields[k]) if k == "meter_interval" and fields[k] is not None else fields[k])
        for k in sorted(fields)
    )
    item_key = tuple(
        (i["text"], i["item_type"], i["unit"], i["min_value"], i["max_value"], bool(i["required"]))
        for i in items
    )
    resource_key = tuple(
        sorted(
            (r["kind"], r["description"], Decimal(r["quantity"]), r["unit"], str(r.get("part_id") or ""))
            for r in resources
        )
    )
    return normalized, item_key, resource_key


def _version_key(version, items, resources):
    fields = {
        "calendar_interval": version.calendar_interval,
        "calendar_unit": version.calendar_unit if version.calendar_interval else CalendarUnit.DAYS,
        "meter_kind": version.meter_kind,
        "meter_interval": version.meter_interval,
        "schedule_mode": version.schedule_mode,
        "tolerance_pct": version.tolerance_pct,
        "estimated_minutes": version.estimated_minutes,
        "instructions": version.instructions,
    }
    return _content_key(
        fields,
        [
            {
                "text": i.text,
                "item_type": i.item_type,
                "unit": i.unit,
                "min_value": i.min_value,
                "max_value": i.max_value,
                "required": i.required,
            }
            for i in items
        ],
        [
            {
                "kind": r.kind,
                "description": r.description,
                "quantity": r.quantity,
                "unit": r.unit,
                "part_id": r.part_id,
            }
            for r in resources
        ],
    )


def _json(value):
    return str(value) if isinstance(value, Decimal) else value


def build_plan_preview(raw: dict, *, membership) -> PlanPreview:
    """Valida el archivo contra los planes y clases de la organización del contexto, sin escribir."""
    require_perm(membership, Perm.PLAN_MANAGE)
    aliases = _choice_aliases()
    classes = {}
    for cls in AssetClass.objects.filter(is_active=True):
        classes.setdefault(cls.code.lower(), cls)
        classes.setdefault(cls.name.lower(), cls)
    existing = {p.code.lower(): p for p in MaintenancePlan.objects.all()}
    drafts = set(
        PlanVersion.objects.filter(status=PlanVersion.Status.DRAFT).values_list("plan_id", flat=True)
    )
    current_ids = [p.current_version_id for p in existing.values() if p.current_version_id]
    current_items, current_resources = {}, {}
    for item in PlanTaskItem.objects.filter(version_id__in=current_ids).order_by("order"):
        current_items.setdefault(item.version_id, []).append(item)
    for resource in PlanResource.objects.filter(version_id__in=current_ids):
        current_resources.setdefault(resource.version_id, []).append(resource)

    parts = _catalog_parts() if any(r.get("part") for r in raw["resources"]) else {}
    codes_in_file = {r.get("code", "").lower() for r in raw["plans"] if r.get("code")}
    orphans, items_by_plan, resources_by_plan = [], {}, {}
    sheet_names = {"items": ITEMS_SHEET, "resources": RESOURCES_SHEET}
    item_errors, resource_errors = {}, {}
    for sheet, target, errors, parser in (
        ("items", items_by_plan, item_errors, _parse_item),
        ("resources", resources_by_plan, resource_errors, _parse_resource),
    ):
        for line in raw[sheet]:
            code = line.get("plan", "").lower()
            if code not in codes_in_file:
                orphans.append(
                    gettext("Hoja %(sheet)s, fila %(row)s: el plan «%(code)s» no está en la hoja Planes.")
                    % {"sheet": sheet_names[sheet], "row": line["_row"], "code": line.get("plan", "")}
                )
                continue
            if not line.get("text" if sheet == "items" else "description"):
                continue
            extra = {"parts": parts} if sheet == "resources" else {}
            parsed = parser(line, aliases, errors.setdefault(code, []), **extra)
            if parsed is not None:
                target.setdefault(code, []).append(parsed)

    rows, seen = [], set()
    for line in raw["plans"]:
        code, name = line.get("code", "").strip(), line.get("name", "").strip()
        row = PlanRow(row=line["_row"], code=code, name=name, action="error")
        rows.append(row)
        key = code.lower()
        if not code:
            row.errors.append(gettext("Falta el código."))
            continue
        if len(code) > 30:
            row.errors.append(gettext("El código supera 30 caracteres."))
        if key in seen:
            row.errors.append(gettext("El código %(code)s se repite en el archivo.") % {"code": code})
        seen.add(key)
        plan = existing.get(key)
        if plan is None and not name:
            row.errors.append(gettext("Falta el nombre."))
        if len(name) > 150:
            row.errors.append(gettext("El nombre supera 150 caracteres."))
        asset_class = None
        if line.get("asset_class"):
            asset_class = classes.get(line["asset_class"].lower())
            if asset_class is None:
                row.errors.append(
                    gettext("La clase «%(value)s» no existe en tu organización.")
                    % {"value": line["asset_class"]}
                )
        fields = _parse_plan(line, aliases, row)
        items = sorted(items_by_plan.get(key, []), key=lambda i: i["_order"])
        items = [{k: v for k, v in i.items() if k != "_order"} for i in items]
        resources = resources_by_plan.get(key, [])
        row.errors += item_errors.get(key, []) + resource_errors.get(key, [])
        if not items:
            row.errors.append(gettext("El plan no tiene ítems en la hoja Checklist."))
        elif len(items) > MAX_ITEMS_PER_PLAN:
            row.errors.append(gettext("Más de %(max)s ítems en el checklist.") % {"max": MAX_ITEMS_PER_PLAN})
        row.item_count = len(items)
        if not row.errors:
            row.frequency_label = str(PlanVersion(**fields).frequency_label)
        if plan is not None:
            row.name = plan.name
            if plan.pk in drafts:
                row.errors.append(
                    gettext("El plan ya tiene un borrador: publícalo o descártalo antes de importar.")
                )
            if (name and name != plan.name) or (asset_class is not None and asset_class != plan.asset_class):
                row.notes.append(
                    gettext("El nombre y la clase de un plan existente no cambian desde el archivo.")
                )
        if row.errors:
            continue
        data = {
            "fields": {k: _json(v) for k, v in fields.items()},
            "items": [{k: _json(v) for k, v in i.items()} for i in items],
            "resources": [{k: _json(v) for k, v in r.items()} for r in resources],
        }
        if plan is None:
            row.action = "create"
            data.update(code=code, name=name, asset_class=str(asset_class.pk) if asset_class else None)
        else:
            version = plan.current_version
            if version is not None and _version_key(
                version, current_items.get(version.id, []), current_resources.get(version.id, [])
            ) == _content_key(fields, items, resources):
                row.action = "unchanged"
                row.notes.append(gettext("Sin cambios respecto de la versión vigente."))
                continue
            row.action = "draft"
            number = (version.number + 1) if version is not None else 1
            row.notes.append(
                gettext("Se crea la versión %(n)s en borrador para revisar y publicar.") % {"n": number}
            )
            data.update(plan=str(plan.pk))
        row.data = data
    return PlanPreview(rows=rows, orphans=orphans)


# --- Guardar ---------------------------------------------------------------------


def _catalog_parts() -> dict:
    """Repuestos activos de la organización por código en minúsculas (una consulta)."""
    from apps.inventory.models import Part

    return {p.code.lower(): p for p in Part.objects.filter(is_active=True)}


def _catalog_parts_ids(session_data) -> list[str]:
    """IDs de repuesto de la sesión que siguen siendo de la organización (una consulta)."""
    from apps.inventory.models import Part

    ids = {r.get("part_id") for row in session_data["rows"] for r in row["data"]["resources"]} - {None}
    if not ids:
        return []
    return [str(pk) for pk in Part.objects.filter(pk__in=ids).values_list("pk", flat=True)]


def _with_part(resource: dict, valid_parts: set) -> dict:
    """El repuesto se conserva solo si sigue en el catálogo de la organización."""
    part_id = resource.pop("part_id", None)
    resource["part_id"] = part_id if part_id in valid_parts else None
    return resource


@dataclass
class PlanImportResult:
    created: int = 0
    drafts: int = 0
    skipped: list[str] = field(default_factory=list)


def _decimals(data: dict, keys) -> dict:
    return {k: (Decimal(v) if k in keys and v is not None else v) for k, v in data.items()}


@transaction.atomic
def apply_plan_import(session_data: dict, *, membership) -> PlanImportResult:
    """Crea los planes nuevos (v1 publicada, por lotes) y los borradores de los existentes."""
    require_perm(membership, Perm.PLAN_MANAGE)
    result = PlanImportResult()
    existing = {p.code.lower(): p for p in MaintenancePlan.objects.all()}
    classes = {str(c.pk): c for c in AssetClass.objects.all()}
    valid_parts = set(_catalog_parts_ids(session_data))
    new_plans = []
    for row in session_data["rows"]:
        data = row["data"]
        fields = _decimals(data["fields"], {"meter_interval"})
        items = [_decimals(i, {"min_value", "max_value"}) for i in data["items"]]
        resources = [_with_part(_decimals(r, {"quantity"}), valid_parts) for r in data["resources"]]
        if row["action"] == "create":
            if data["code"].lower() in existing:  # se creó entre la vista previa y la confirmación
                result.skipped.append(data["code"])
                continue
            new_plans.append(
                {
                    "code": data["code"],
                    "name": data["name"],
                    "asset_class": classes.get(data["asset_class"]) if data["asset_class"] else None,
                    "fields": fields,
                    "items": items,
                    "resources": resources,
                }
            )
            continue
        plan = MaintenancePlan.objects.filter(pk=data["plan"]).first()
        if plan is None or plan.versions.filter(status=PlanVersion.Status.DRAFT).exists():
            result.skipped.append(row["code"])
            continue
        create_draft_with_content(
            plan=plan, fields=fields, items=items, resources=resources, by=membership.user
        )
        result.drafts += 1
    if new_plans:
        created = bulk_create_published_plans(
            organization=membership.organization,
            plans=new_plans,
            by=membership.user,
            change_note=gettext("Importado desde Excel"),
        )
        for plan in created:
            audit.record(
                action=AuditEvent.Action.CREATE,
                target=plan,
                membership=membership,
                changes={"code": [None, plan.code], "source": [None, "excel"]},
            )
        result.created = len(created)
    return result
