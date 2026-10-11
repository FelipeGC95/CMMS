"""Importación de repuestos y stock desde Excel o CSV (INV-02).

Mismo patrón que la importación de activos (AST-08): `parse_file` + `build_preview`
validan sin escribir; `apply_import` guarda en una sola transacción. A diferencia
de activos, es **todo o nada**: con una fila con error no se importa ninguna.

Idempotente por código de repuesto (crea o actualiza; las celdas vacías no borran
datos). La columna `cantidad` es un **conteo** en esa bodega (o en la por
defecto): se registra un ajuste con motivo "import" por la diferencia, y si no
hay diferencia no se crea movimiento. Con costo y existencia total ≤ 0, el costo
fija el promedio (lo resuelve `count_stock`, INV-08).
"""

import csv
import datetime
import io
import re
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.translation import gettext, gettext_lazy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from . import services
from .models import Part, StockItem, StockMovement
from .selectors import can_see_warehouse, default_warehouse, visible_warehouses

MAX_ROWS = 2000
MAX_FILE_BYTES = 5 * 1024 * 1024
SHEET_NAME = gettext_lazy("Repuestos")

# (clave interna, encabezado de la plantilla, obligatorio, ayuda)
COLUMNS = [
    (
        "code",
        gettext_lazy("Código"),
        True,
        gettext_lazy("Único por repuesto. Si ya existe, la fila actualiza ese repuesto."),
    ),
    ("name", gettext_lazy("Nombre"), True, gettext_lazy("Obligatorio solo para repuestos nuevos.")),
    ("unit", gettext_lazy("Unidad"), False, gettext_lazy("Ej.: un, l, kg, m. Por defecto «un».")),
    ("category", gettext_lazy("Categoría"), False, gettext_lazy("Ej.: Rodamientos, Filtros, Correas.")),
    ("manufacturer_ref", gettext_lazy("N° de parte"), False, gettext_lazy("Número de parte del fabricante.")),
    ("manufacturer", gettext_lazy("Fabricante"), False, ""),
    ("supplier", gettext_lazy("Proveedor"), False, ""),
    (
        "warehouse",
        gettext_lazy("Bodega"),
        False,
        gettext_lazy("Código o nombre de la bodega. Vacío: la bodega por defecto."),
    ),
    (
        "quantity",
        gettext_lazy("Cantidad"),
        False,
        gettext_lazy("Lo que contaste hoy en esa bodega. Se ajusta el stock a esta cantidad."),
    ),
    (
        "unit_cost",
        gettext_lazy("Costo unitario"),
        False,
        gettext_lazy("Fija el costo promedio si el repuesto no tenía stock."),
    ),
    ("min_quantity", gettext_lazy("Stock mínimo"), False, gettext_lazy("Bajo esta cantidad se avisa.")),
    ("target_quantity", gettext_lazy("Stock objetivo"), False, gettext_lazy("Hasta dónde reponer.")),
    ("bin_location", gettext_lazy("Ubicación en bodega"), False, gettext_lazy("Ej.: Estante A-3.")),
]
REQUIRED_KEYS = {"code", "name"}
# Encabezados reconocidos (normalizados: sin tildes, minúsculas y "_").
HEADER_ALIASES = {
    "codigo": "code",
    "code": "code",
    "nombre": "name",
    "name": "name",
    "descripcion": "name",
    "unidad": "unit",
    "unit": "unit",
    "categoria": "category",
    "familia": "category",
    "n_de_parte": "manufacturer_ref",
    "no_de_parte": "manufacturer_ref",
    "numero_de_parte": "manufacturer_ref",
    "n_parte": "manufacturer_ref",
    "part_number": "manufacturer_ref",
    "fabricante": "manufacturer",
    "marca": "manufacturer",
    "proveedor": "supplier",
    "bodega": "warehouse",
    "cantidad": "quantity",
    "stock": "quantity",
    "existencia": "quantity",
    "costo_unitario": "unit_cost",
    "costo": "unit_cost",
    "precio_unitario": "unit_cost",
    "stock_minimo": "min_quantity",
    "minimo": "min_quantity",
    "stock_objetivo": "target_quantity",
    "objetivo": "target_quantity",
    "maximo": "target_quantity",
    "ubicacion_en_bodega": "bin_location",
    "ubicacion": "bin_location",
}
TEXT_LIMITS = {
    "code": 40,
    "name": 200,
    "unit": 20,
    "category": 80,
    "manufacturer_ref": 80,
    "manufacturer": 120,
    "supplier": 150,
    "bin_location": 40,
}
NUMBER_KEYS = ("quantity", "unit_cost", "min_quantity", "target_quantity")


class ImportFileError(ValidationError):
    """El archivo no se puede leer (no las filas: esas se informan una a una)."""


def _normalize(text) -> str:
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _column_label(key) -> str:
    return str(next(label for k, label, *_rest in COLUMNS if k == key))


# --- Lectura del archivo ----------------------------------------------------


def _xlsx_rows(uploaded_file):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ImportFileError(
            gettext("No se pudo leer el archivo. Usa un Excel .xlsx o un CSV (idealmente la plantilla).")
        ) from exc
    sheet_name = str(SHEET_NAME)
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.worksheets[0]
    rows = [list(values) for values in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def _csv_rows(uploaded_file):
    raw = uploaded_file.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:  # CSV guardado por Excel en Windows
        text = raw.decode("cp1252", errors="replace")
    first_line = text.splitlines()[0] if text else ""
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def parse_file(uploaded_file) -> list[dict]:
    """Lee un .xlsx (primera hoja o «Repuestos») o un .csv (coma o punto y coma). Filas como texto."""
    if uploaded_file.size > MAX_FILE_BYTES:
        raise ImportFileError(gettext("El archivo supera 5 MB."))
    name = (uploaded_file.name or "").lower()
    rows = _csv_rows(uploaded_file) if name.endswith(".csv") else _xlsx_rows(uploaded_file)
    header = rows[0] if rows else []
    keys = [HEADER_ALIASES.get(_normalize(h)) if h not in (None, "") else None for h in header]
    missing = [_column_label(key) for key in ("code", "name") if key not in keys]
    if missing:
        raise ImportFileError(
            gettext("Faltan columnas obligatorias: %(columns)s.") % {"columns": ", ".join(missing)}
        )
    parsed = []
    for number, values in enumerate(rows[1:], start=2):
        row = {}
        for key, value in zip(keys, values, strict=False):
            if key and key not in row:
                row[key] = _cell_text(value)
        if not any(row.values()):
            continue
        if len(parsed) >= MAX_ROWS:
            raise ImportFileError(
                gettext("El archivo tiene más de %(max)s filas; divídelo en partes.") % {"max": MAX_ROWS}
            )
        row["_row"] = number
        parsed.append(row)
    if not parsed:
        raise ImportFileError(gettext("El archivo no tiene filas con datos."))
    return parsed


# --- Validación (sin escribir) ---------------------------------------------


@dataclass
class RowResult:
    row: int
    code: str
    name: str
    action: str  # create | update | error
    warehouse: str = ""
    warehouse_id: str = ""
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    data: dict = field(default_factory=dict)


@dataclass
class Preview:
    rows: list[RowResult]

    @property
    def has_errors(self) -> bool:
        return any(row.action == "error" for row in self.rows)

    def count(self, action) -> int:
        """Repuestos (códigos distintos) que se crean o actualizan; filas con error."""
        if action == "error":
            return sum(1 for row in self.rows if row.action == "error")
        return len({row.code.lower() for row in self.rows if row.action == action})

    @property
    def counts(self) -> int:
        """Filas con cantidad (conteos que se registran)."""
        return sum(1 for row in self.rows if row.data.get("quantity", "") != "")

    def to_session(self) -> dict:
        return {"rows": [asdict(row) for row in self.rows]}


def _number(text):
    """Decimal ≥ 0 con coma o punto decimal; None si está vacío; InvalidOperation si no es válido."""
    if text == "":
        return None
    value = Decimal(text.replace(" ", "").replace(",", "."))
    if not value.is_finite() or value < 0:
        raise InvalidOperation
    return value


def _warehouse_index(warehouses) -> dict:
    index = {}
    for warehouse in warehouses:
        index.setdefault(warehouse.code.lower(), warehouse)
        index.setdefault(warehouse.name.lower(), warehouse)
    return index


def _format(value: Decimal) -> str:
    text = f"{value:f}"
    return (text.rstrip("0").rstrip(".") if "." in text else text).replace(".", ",") or "0"


def build_preview(raw_rows: list[dict], *, membership) -> Preview:
    """Valida cada fila y describe lo que hará (crear, actualizar, conteo), sin escribir nada."""
    warehouses = _warehouse_index(visible_warehouses(membership=membership))
    default = default_warehouse()
    existing = {part.code.lower(): part for part in Part.objects.all()}
    items = {
        (item.part_id, item.warehouse_id): item
        for item in StockItem.objects.filter(part__in=list(existing.values()))
    }
    seen, named = set(), set()
    results = []
    for raw in raw_rows:
        code, name = raw.get("code", ""), raw.get("name", "")
        result = RowResult(row=raw["_row"], code=code, name=name, action="create")
        errors, notes = result.errors, result.notes
        part = existing.get(code.lower()) if code else None
        if part is not None:
            result.action = "update"
            result.name = name or part.name
        if not code:
            errors.append(gettext("Falta el código."))
        # Un repuesto nuevo puede repetirse en otra bodega: basta el nombre en su primera fila.
        if not name and part is None and code.lower() not in named:
            errors.append(gettext("Falta el nombre."))
        if name and code:
            named.add(code.lower())
        for key, limit in TEXT_LIMITS.items():
            if len(raw.get(key, "")) > limit:
                errors.append(
                    gettext("%(column)s: máximo %(n)s caracteres.")
                    % {"column": _column_label(key), "n": limit}
                )

        # Bodega: la indicada (visible y activa) o la por defecto (se crea sola si no hay).
        warehouse_value = raw.get("warehouse", "")
        if warehouse_value:
            warehouse = warehouses.get(warehouse_value.lower())
            if warehouse is None:
                errors.append(gettext("No existe la bodega «%(name)s».") % {"name": warehouse_value})
        else:
            warehouse = default
            if warehouse is not None and not can_see_warehouse(membership, warehouse):
                errors.append(gettext("Indica una bodega de los sitios a los que tienes acceso."))
        if warehouse is not None:
            result.warehouse, result.warehouse_id = warehouse.code, str(warehouse.pk)
        else:
            result.warehouse = warehouse_value or services.DEFAULT_WAREHOUSE_CODE
        key = (code.lower(), result.warehouse_id or "default")
        if code and key in seen:
            errors.append(gettext("Código repetido en el archivo para la misma bodega."))
        seen.add(key)

        numbers = {}
        for number_key in NUMBER_KEYS:
            try:
                numbers[number_key] = _number(raw.get(number_key, ""))
            except InvalidOperation:
                errors.append(
                    gettext("%(column)s: debe ser un número mayor o igual a cero.")
                    % {"column": _column_label(number_key)}
                )
                numbers[number_key] = None

        item = items.get((part.pk, warehouse.pk)) if part is not None and warehouse is not None else None
        minimum = numbers["min_quantity"]
        target = numbers["target_quantity"]
        effective_min = minimum if minimum is not None else (item.min_quantity if item else Decimal(0))
        if target is not None and target < effective_min:
            errors.append(gettext("El stock objetivo no puede ser menor que el mínimo."))

        quantity, cost = numbers["quantity"], numbers["unit_cost"]
        if quantity is not None:
            previous = item.on_hand if item else Decimal(0)
            if quantity == previous:
                notes.append(gettext("Stock sin cambios (%(n)s).") % {"n": _format(quantity)})
            else:
                notes.append(
                    gettext("Conteo: %(before)s → %(after)s.")
                    % {"before": _format(previous), "after": _format(quantity)}
                )
        elif cost is not None:
            notes.append(gettext("Sin cantidad: el costo unitario no se usa."))

        if errors:
            result.action = "error"
        result.data = {k: v for k, v in raw.items() if k != "_row"}
        results.append(result)
    return Preview(results)


# --- Aplicación -------------------------------------------------------------


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    adjusted: int = 0
    unchanged: int = 0


PART_KEYS = ("name", "unit", "category", "manufacturer_ref", "manufacturer", "supplier")


@transaction.atomic
def apply_import(session_data: dict, *, membership) -> ImportResult:
    """Guarda todas las filas de la vista previa en una transacción (todo o nada, INV-02 CA5).

    Cada repuesto y cada conteo pasa por los servicios (auditoría, costo promedio y
    alerta de mínimo); un error en cualquier fila revierte la importación completa.
    """
    rows = session_data["rows"]
    if any(row["action"] == "error" for row in rows):
        raise ValidationError(gettext("El archivo tiene filas con errores: corrígelas y vuelve a subirlo."))
    result = ImportResult()
    parts = {part.code.lower(): part for part in Part.objects.all()}
    touched = set()
    for row in rows:
        data = row["data"]
        code = data["code"]
        part = parts.get(code.lower())
        fields = {key: data[key] for key in PART_KEYS if data.get(key, "")}
        if part is None:
            part = services.create_part(membership=membership, code=code, **fields)
            parts[code.lower()] = part
            result.created += 1
            touched.add(part.pk)
        else:
            changed = {key: value for key, value in fields.items() if getattr(part, key) != value}
            if changed:
                part = services.update_part(membership=membership, part=part, **changed)
            if part.pk not in touched:
                result.updated += 1
                touched.add(part.pk)

        warehouse = row["warehouse_id"] or None
        if data.get("quantity", "") != "":
            movement = services.count_stock(
                membership=membership,
                part=part,
                counted_quantity=data["quantity"],
                warehouse=warehouse,
                unit_cost=data.get("unit_cost") or None,
                reason=StockMovement.Reason.IMPORT,
                note=gettext("Importación, fila %(row)s") % {"row": row["row"]},
            )
            if movement is None:
                result.unchanged += 1
            else:
                result.adjusted += 1

        level_keys = ("min_quantity", "target_quantity", "bin_location")
        if any(data.get(key, "") != "" for key in level_keys):
            current = StockItem.objects.filter(
                part=part, warehouse=warehouse or services.ensure_default_warehouse()
            ).first()
            services.set_stock_levels(
                membership=membership,
                part=part,
                warehouse=warehouse,
                min_quantity=data.get("min_quantity") or (current.min_quantity if current else 0),
                target_quantity=data.get("target_quantity") or (current.target_quantity if current else None),
                bin_location=data.get("bin_location") or (current.bin_location if current else ""),
            )
    return result


# --- Plantilla --------------------------------------------------------------


def build_template(*, membership) -> bytes:
    """Plantilla .xlsx con encabezados, un ejemplo, instrucciones y las bodegas existentes."""
    wb = Workbook()
    ws = wb.active
    ws.title = str(SHEET_NAME)
    header_fill = PatternFill("solid", fgColor="0B2E4F")
    for col, (_key, label, _required, _help) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=str(label))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(label)) + 6)
    default = default_warehouse()
    example = [
        "ROD-6205",
        gettext("Rodamiento 6205 2RS"),
        "un",
        gettext("Rodamientos"),
        "6205-2RSH",
        "SKF",
        "",
        default.code if default else services.DEFAULT_WAREHOUSE_CODE,
        12,
        4500,
        4,
        10,
        "A-3",
    ]
    for col, value in enumerate(example, start=1):
        ws.cell(row=2, column=col, value=value).font = Font(italic=True, color="475569")
    ws.freeze_panes = "A2"

    help_sheet = wb.create_sheet(gettext("Instrucciones"))
    help_sheet.column_dimensions["A"].width = 22
    help_sheet.column_dimensions["B"].width = 90
    help_sheet.append([gettext("Columna"), gettext("Cómo llenarla")])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for key, label, _required, help_text in COLUMNS:
        column = (
            gettext("%(column)s (obligatoria)") % {"column": label} if key in REQUIRED_KEYS else str(label)
        )
        help_sheet.append([column, str(help_text)])
    help_sheet.append([])
    help_sheet.append(["", gettext("Borra la fila de ejemplo (en cursiva) antes de subir el archivo.")])
    help_sheet.append(
        [
            "",
            gettext(
                "Si hay un error en cualquier fila no se importa nada: la vista previa te dice qué corregir."
            ),
        ]
    )
    help_sheet.append(
        [
            "",
            gettext("Máximo %(max)s filas por archivo. También se acepta CSV con las mismas columnas.")
            % {"max": MAX_ROWS},
        ]
    )

    reference = wb.create_sheet(gettext("Bodegas"))
    reference.append([gettext("Código"), gettext("Nombre")])
    for cell in reference[1]:
        cell.font = Font(bold=True)
    for warehouse in visible_warehouses(membership=membership):
        reference.append([warehouse.code, warehouse.name])
    reference.column_dimensions["A"].width = 16
    reference.column_dimensions["B"].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_filename() -> str:
    return f"plantilla-repuestos-{datetime.date.today():%Y%m%d}.xlsx"
