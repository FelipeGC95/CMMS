"""Exportación de listas a Excel (KPI-03) con openpyxl en modo de solo escritura.

Cada exportación respeta los filtros de su lista y el alcance por sitio: recibe
la consulta ya filtrada por el selector de la lista.
"""

import io

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from apps.ui.templatetags.ui import work_order_status_label

# Tope de filas por archivo: por sobre esto se pide filtrar la lista (el archivo se arma en memoria).
MAX_EXPORT_ROWS = 10_000

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _local(value):
    """Excel no admite zonas horarias: fecha y hora locales, sin zona."""
    return timezone.localtime(value).replace(tzinfo=None) if value else None


def _workbook(title, headers, rows) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=title[:31])
    sheet.freeze_panes = "A2"
    bold = Font(bold=True)
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=str(header))
        cell.font = bold
        header_cells.append(cell)
    sheet.append(header_cells)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def assets_xlsx(assets) -> bytes:
    """`assets`: consulta de la lista de activos (ya filtrada y con alcance)."""
    headers = [
        _("Código"),
        _("Nombre"),
        _("Clase"),
        _("Ubicación"),
        _("Activo padre"),
        _("Estado"),
        _("Criticidad"),
        _("Fabricante"),
        _("Modelo"),
        _("N.º de serie"),
        _("Año"),
        _("Puesta en servicio"),
    ]
    rows = (
        (
            a.code,
            a.name,
            a.asset_class.name,
            a.location.name if a.location else "",
            a.parent.code if a.parent else "",
            a.get_status_display(),
            a.criticality,
            a.manufacturer,
            a.model,
            a.serial_number,
            a.year,
            a.in_service_on,
        )
        for a in assets.select_related("asset_class", "location", "parent").order_by("code").iterator()
    )
    return _workbook(_("Activos"), headers, rows)


def work_orders_xlsx(work_orders) -> bytes:
    """`work_orders`: consulta de la lista de OT (ya filtrada y con alcance), en su orden."""
    headers = [
        _("OT"),
        _("Tipo"),
        _("Estado"),
        _("Prioridad"),
        _("Título"),
        _("Activo o ubicación"),
        _("Responsable"),
        _("Fecha límite"),
        _("Creada"),
        _("Completada"),
        _("Cerrada"),
        _("Costo total"),
    ]
    rows = (
        (
            wo.code,
            wo.get_work_type_display(),
            str(work_order_status_label(wo.status)),
            wo.get_priority_display(),
            wo.title,
            str(wo.asset.code if wo.asset_id else wo.location.name if wo.location_id else ""),
            (wo.assigned_to.full_name or wo.assigned_to.email) if wo.assigned_to_id else "",
            wo.due_on,
            _local(wo.created_at),
            _local(wo.completed_at),
            _local(wo.closed_at),
            wo.total_cost,
        )
        for wo in work_orders.select_related("asset", "location", "assigned_to").iterator()
    )
    return _workbook(_("Órdenes de trabajo"), headers, rows)


def backlog_xlsx(work_orders) -> bytes:
    """`work_orders`: backlog filtrado (REQ-04), anotado con `score` y en su orden."""
    headers = [
        _("Puntaje"),
        _("OT"),
        _("Tipo"),
        _("Estado"),
        _("Prioridad"),
        _("Criticidad"),
        _("Título"),
        _("Activo o ubicación"),
        _("Responsable"),
        _("Fecha límite"),
        _("Programada para"),
        _("Estimación (h)"),
    ]
    rows = (
        (
            wo.score,
            wo.code,
            wo.get_work_type_display(),
            str(work_order_status_label(wo.status)),
            wo.get_priority_display(),
            wo.asset.criticality if wo.asset_id else "",
            wo.title,
            str(wo.asset.code if wo.asset_id else wo.location.name if wo.location_id else ""),
            (wo.assigned_to.full_name or wo.assigned_to.email) if wo.assigned_to_id else "",
            wo.due_on,
            wo.scheduled_for,
            round(wo.estimated_minutes / 60, 1) if wo.estimated_minutes else None,
        )
        for wo in work_orders.iterator()
    )
    return _workbook(_("Backlog"), headers, rows)


def _sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title=title[:31])
    sheet.freeze_panes = "A2"
    bold = Font(bold=True)
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=str(header))
        cell.font = bold
        header_cells.append(cell)
    sheet.append(header_cells)
    for row in rows:
        sheet.append(row)


def _round(value, digits=1):
    return round(float(value), digits) if value is not None else None


def extended_kpis_xlsx(*, filter_rows, backlog, mix, costs, pareto) -> bytes:
    """Indicadores ampliados (KPI-02) a Excel: una hoja con los filtros y una por indicador.

    Recibe los resultados ya calculados por los selectores de `extended` con los
    filtros y el alcance de la pantalla; los costos por activo van completos (sin tope).
    """
    workbook = Workbook(write_only=True)
    _sheet(workbook, _("Filtros"), [_("Filtro"), _("Valor")], [(str(k), str(v)) for k, v in filter_rows])
    _sheet(
        workbook,
        _("Backlog en semanas"),
        [
            _("Al día"),
            _("OT pendientes"),
            _("Horas estimadas"),
            _("OT sin estimación"),
            _("Capacidad semanal (h)"),
            _("Backlog (semanas)"),
        ],
        (
            (
                p.day,
                p.count,
                _round(p.estimated_hours),
                p.unestimated,
                _round(backlog.weekly_capacity_hours),
                _round(p.weeks),
            )
            for p in backlog.points
        ),
    )
    _sheet(
        workbook,
        _("Preventivo-correctivo"),
        [
            _("Mes"),
            _("OT planificadas"),
            _("OT correctivas"),
            _("OT de mejora"),
            _("Horas planificadas"),
            _("Horas correctivas"),
            _("Horas de mejora"),
            _("% planificado (horas)"),
            _("% planificado (cantidad)"),
        ],
        (
            (
                row.month.strftime("%Y-%m") if row.month else str(_("Total")),
                row.planned_count,
                row.corrective_count,
                row.other_count,
                _round(row.planned_hours),
                _round(row.corrective_hours),
                _round(row.other_hours),
                _round(row.planned_pct),
                _round(row.planned_pct_count),
            )
            for row in [*mix.months, mix.total]
        ),
    )
    money = [_("Mano de obra"), _("Repuestos"), _("Otros costos"), _("Costo total")]
    _sheet(
        workbook,
        _("Costos por activo"),
        [_("Código"), _("Activo"), _("Clase"), _("OT cerradas"), *money],
        (
            (
                r["asset__code"],
                r["asset__name"],
                r["asset__asset_class__name"],
                r["work_orders"],
                r["labor"],
                r["parts"],
                r["other"],
                r["total"],
            )
            for r in costs.by_asset.iterator()
        ),
    )
    loc, totals = costs.location_totals, costs.totals
    _sheet(
        workbook,
        _("Costos por clase"),
        [_("Clase"), _("Activos"), _("OT cerradas"), *money],
        [
            *(
                (
                    r["asset__asset_class__name"],
                    r["assets"],
                    r["work_orders"],
                    r["labor"],
                    r["parts"],
                    r["other"],
                    r["total"],
                )
                for r in costs.by_class
            ),
            (
                str(_("OT sobre ubicaciones (sin activo)")),
                None,
                loc["work_orders"],
                loc["labor"],
                loc["parts"],
                loc["other"],
                loc["total"],
            ),
            (
                str(_("Total")),
                None,
                totals["work_orders"],
                totals["labor"],
                totals["parts"],
                totals["other"],
                totals["total"],
            ),
        ],
    )
    code_label = _("Modo de falla") if pareto.code_kind == "mode" else _("Causa")
    share_label = _("% de fallas") if pareto.metric == "count" else _("% de la detención")
    _sheet(
        workbook,
        _("Pareto de fallas"),
        [
            _("Código"),
            code_label,
            _("Fallas"),
            _("Detención (h)"),
            share_label,
            _("% acumulado"),
            _("Dentro del 80 %"),
        ],
        [
            *(
                (
                    r.code,
                    r.name,
                    r.count,
                    _round(r.downtime_hours),
                    _round(r.share_pct),
                    _round(r.cumulative_pct),
                    str(_("Sí")) if r.vital else str(_("No")),
                )
                for r in pareto.rows
            ),
            (
                None,
                str(_("Sin código (fuera del Pareto)")),
                pareto.uncoded_count,
                _round(pareto.uncoded_downtime_hours),
                None,
                None,
                None,
            ),
        ],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Inventario (INV-14 CA5) ------------------------------------------------------
# Los costos van solo con `inventory.cost_view` (Q3): sin él, las columnas no existen.


def _stock_state(part) -> str:
    if not part.is_active:
        return str(_("Inactivo"))
    if part.negative:
        return str(_("Negativo · regularizar"))
    if part.below_min:
        return str(_("Bajo mínimo"))
    if part.on_hand_total == 0:
        return str(_("Sin stock"))
    return str(_("Con stock"))


def stock_xlsx(parts, *, show_costs: bool) -> bytes:
    """`parts`: consulta de `inventory.selectors.stock_list` (filtros y alcance de la pantalla)."""
    headers = [
        _("Código"),
        _("Nombre"),
        _("Unidad"),
        _("Categoría"),
        _("N° de parte"),
        _("Fabricante"),
        _("Proveedor"),
        _("En existencia"),
        _("Reservado"),
        _("Disponible"),
        _("Estado"),
    ]
    if show_costs:
        headers += [_("Costo promedio"), _("Valor en existencia")]

    def row(part):
        values = [
            part.code,
            part.name,
            part.unit,
            part.category,
            part.manufacturer_ref,
            part.manufacturer,
            part.supplier,
            part.on_hand_total,
            part.reserved_total,
            part.available_total,
            _stock_state(part),
        ]
        if show_costs:
            values += [part.avg_cost, round(max(part.on_hand_total, 0) * part.avg_cost, 2)]
        return values

    return _workbook(_("Stock"), headers, (row(part) for part in parts.iterator()))


def movements_xlsx(movements, *, show_costs: bool) -> bytes:
    """`movements`: consulta de `inventory.selectors.movement_list` con los filtros aplicados."""
    headers = [
        _("Fecha"),
        _("Código"),
        _("Repuesto"),
        _("Tipo"),
        _("Motivo"),
        _("Bodega"),
        _("Cantidad"),
        _("Unidad"),
        _("Saldo resultante"),
    ]
    if show_costs:
        headers += [_("Costo unitario"), _("Costo total")]
    headers += [_("OT"), _("Documento"), _("Proveedor"), _("Nota"), _("Registrado por"), _("Registrado el")]

    def row(m):
        values = [
            _local(m.occurred_at),
            m.part.code,
            m.part.name,
            m.get_kind_display(),
            m.get_reason_display() if m.reason else "",
            m.warehouse.name,
            m.quantity,
            m.part.unit,
            m.on_hand_after,
        ]
        if show_costs:
            values += [m.unit_cost, m.total_cost]
        values += [
            m.work_order.code if m.work_order_id else "",
            m.reference,
            m.supplier,
            m.note,
            (m.created_by.full_name or m.created_by.email) if m.created_by_id else "",
            _local(m.created_at),
        ]
        return values

    return _workbook(_("Movimientos"), headers, (row(m) for m in movements.iterator()))


def reorder_xlsx(items, *, show_costs: bool) -> bytes:
    """`items`: consulta de `inventory.selectors.reorder_list` (por bodega, prioridad y código)."""
    headers = [
        _("Bodega"),
        _("Código"),
        _("Repuesto"),
        _("Unidad"),
        _("Proveedor"),
        _("N° de parte"),
        _("En existencia"),
        _("Reservado"),
        _("Disponible"),
        _("Stock mínimo"),
        _("Stock objetivo"),
        _("Faltante de reservas"),
        _("Sugerido"),
        _("Estado"),
    ]
    if show_costs:
        headers += [_("Último costo"), _("Costo estimado")]
    states = {0: _("Negativo · regularizar"), 1: _("Reservas sin cubrir"), 2: _("Bajo mínimo")}

    def row(item):
        part = item.part
        values = [
            item.warehouse.name,
            part.code,
            part.name,
            part.unit,
            part.supplier,
            part.manufacturer_ref,
            item.on_hand,
            item.reserved,
            item.available,
            item.min_quantity,
            item.target_quantity,
            item.reservation_gap,
            item.suggested,
            str(states[item.priority]),
        ]
        if show_costs:
            cost = part.last_cost if part.last_cost is not None else part.avg_cost
            values += [cost, round(item.suggested * cost, 2)]
        return values

    return _workbook(_("Reposición"), headers, (row(item) for item in items.iterator()))
