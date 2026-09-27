"""KPI-03: listas de activos y OT a Excel, con los filtros de la pantalla y un tope de filas."""

import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.assets.models import Asset
from apps.core.tenancy import tenant_context
from apps.work.services import create_work_order
from tests.conftest import login


def xlsx_rows(response):
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert "attachment;" in response["Content-Disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


@pytest.fixture
def data(tenant_a):
    with tenant_context(tenant_a.org):
        cls = tenant_a.asset.asset_class
        Asset.objects.create(code="CMP-02", name="Compresor 2", asset_class=cls, status="stopped")
        Asset.objects.create(code="BOM-01", name="Bomba Ñandú", asset_class=cls)
        create_work_order(work_type="corrective", title="Fuga", asset=tenant_a.asset)
        create_work_order(work_type="preventive", title="PM", asset=tenant_a.asset)
        create_work_order(work_type="corrective", title="Hecha", asset=tenant_a.asset, status="closed")


def test_asset_export_headers_and_rows(client, tenant_a, data):
    login(client, tenant_a.admin)
    rows = xlsx_rows(client.get(reverse("asset_export_xlsx")))
    assert rows[0][:3] == ["Código", "Nombre", "Clase"]
    assert [r[0] for r in rows[1:]] == ["BOM-01", "CMP-01", "CMP-02"]
    assert rows[1][1] == "Bomba Ñandú"


@pytest.mark.parametrize(
    ("params", "codes"),
    [({"q": "cmp"}, ["CMP-01", "CMP-02"]), ({"status": "stopped"}, ["CMP-02"]), ({"q": "zzz"}, [])],
)
def test_asset_export_honors_list_filters(client, tenant_a, data, params, codes):
    login(client, tenant_a.admin)
    rows = xlsx_rows(client.get(reverse("asset_export_xlsx"), params))
    assert [r[0] for r in rows[1:]] == codes


def test_work_order_export_honors_scope_and_type(client, tenant_a, data):
    login(client, tenant_a.admin)
    url = reverse("work_order_export_xlsx")
    rows = xlsx_rows(client.get(url))
    assert rows[0][:5] == ["OT", "Tipo", "Estado", "Prioridad", "Título"]
    assert sorted(r[4] for r in rows[1:]) == ["Fuga", "PM"]  # abiertas por defecto
    assert [r[4] for r in xlsx_rows(client.get(url, {"scope": "done"}))[1:]] == ["Hecha"]
    typed = xlsx_rows(client.get(url, {"type": "preventive"}))
    assert [r[4] for r in typed[1:]] == ["PM"]
    assert typed[1][5] == "CMP-01"


def test_export_over_the_cap_returns_to_the_list(client, tenant_a, data, monkeypatch):
    monkeypatch.setattr("apps.reporting.views.MAX_EXPORT_ROWS", 2)
    login(client, tenant_a.admin)
    response = client.get(reverse("asset_export_xlsx"), {"q": "o"}, follow=True)
    assert response.redirect_chain[-1][0] == f"{reverse('asset_list')}?q=o"
    assert "Filtra la lista" in response.content.decode()


def test_list_pages_link_to_the_export_with_current_filters(client, tenant_a):
    login(client, tenant_a.admin)
    body = client.get(reverse("asset_list"), {"status": "stopped"}).content.decode()
    assert f"{reverse('asset_export_xlsx')}?status=stopped" in body
    body = client.get(reverse("work_order_list"), {"scope": "done"}).content.decode()
    assert f"{reverse('work_order_export_xlsx')}?scope=done" in body
