"""KPI-02: la organización B no ve datos de A en ningún indicador ampliado ni en su Excel."""

import datetime
import io

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.tenancy import tenant_context
from apps.organizations.models import Location
from apps.work.models import FailureCode, FailureEvent, LaborEntry, WorkOrder
from apps.work.services import create_work_order
from tests.conftest import login


@pytest.fixture
def data_in_a(tenant_a):
    """En A: una OT pendiente, una preventiva cerrada con costo y horas, y una falla con código."""
    yesterday = timezone.now() - datetime.timedelta(days=1)
    with tenant_context(tenant_a.org):
        asset = tenant_a.asset
        create_work_order(work_type="corrective", title="Pendiente A", asset=asset, estimated_minutes=600)
        done = create_work_order(work_type="preventive", title="Cerrada A", asset=asset)
        LaborEntry.objects.create(work_order=done, user=tenant_a.tech.user, minutes=120)
        WorkOrder.objects.filter(pk=done.pk).update(
            status="closed", completed_at=yesterday, closed_at=yesterday, total_cost=9999, labor_cost=9999
        )
        mode = FailureCode.objects.create(kind="mode", code="SOLOA", name="Modo de A")
        FailureEvent.objects.create(
            asset=asset, impact="total_stop", failure_mode=mode, started_at=yesterday, ended_at=timezone.now()
        )
        site = Location.objects.create(code="SITIO-A", name="Sitio A")
    return {"site": site}


def test_report_of_b_has_nothing_from_a(client, tenant_a, tenant_b, data_in_a):
    login(client, tenant_b.admin)
    context = client.get(reverse("kpi_report")).context
    assert all(p.count == 0 for p in context["backlog"].points)
    assert context["mix"].total.planned_count == 0
    assert context["costs"].totals["total"] == 0 and context["cost_rows"] == []
    assert context["pareto"].rows == [] and context["pareto"].uncoded_count == 0


def test_report_of_a_sees_its_own_data(client, tenant_a, data_in_a):
    login(client, tenant_a.admin)
    context = client.get(reverse("kpi_report")).context
    assert context["backlog"].current.count == 1
    assert context["mix"].total.planned_count == 1
    assert context["costs"].totals["total"] == 9999
    assert [r.code for r in context["pareto"].rows] == ["SOLOA"]


def test_excel_of_b_has_nothing_from_a(client, tenant_a, tenant_b, data_in_a):
    login(client, tenant_b.admin)
    response = client.get(reverse("kpi_report_export_xlsx"))
    book = load_workbook(io.BytesIO(response.content))
    assert len(list(book["Costos por activo"].iter_rows())) == 1  # solo encabezados
    pareto = [r for r in book["Pareto de fallas"].iter_rows(values_only=True)]
    assert all("SOLOA" not in (cell or "") for row in pareto for cell in row if isinstance(cell, str))
    assert all(row[1] == 0 for row in list(book["Backlog en semanas"].iter_rows(values_only=True))[1:])


def test_foreign_location_filter_is_rejected(client, tenant_a, tenant_b, data_in_a):
    login(client, tenant_b.admin)
    response = client.get(reverse("kpi_report"), {"location": data_in_a["site"].pk})
    assert "location" in response.context["filter_form"].errors
    excel = client.get(reverse("kpi_report_export_xlsx"), {"location": data_in_a["site"].pk})
    filters = list(load_workbook(io.BytesIO(excel.content))["Filtros"].iter_rows(values_only=True))
    assert all("Sitio A" not in str(row[1]) for row in filters)


def test_foreign_class_and_asset_filters_are_rejected(client, tenant_a, tenant_b, data_in_a):
    login(client, tenant_b.admin)
    with tenant_context(tenant_a.org):
        from apps.assets.models import Asset

        Asset.objects.create(code="SOLO-A", name="x", asset_class=tenant_a.asset.asset_class)
    response = client.get(
        reverse("kpi_report"), {"asset_class": tenant_a.asset.asset_class_id, "asset": "SOLO-A"}
    )
    errors = response.context["filter_form"].errors
    assert "asset_class" in errors and "asset" in errors
