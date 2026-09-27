"""Planes por Excel (PM-08): exportar e importar nunca ven ni referencian datos de otra organización."""

import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.core.tenancy import tenant_context
from apps.maintenance.models import MaintenancePlan, PlanVersion
from tests.conftest import login


def _upload(plans, items):
    wb = Workbook()
    ws = wb.active
    ws.title = "Planes"
    ws.append(["Código", "Nombre", "Clase", "Cada (calendario)", "Unidad de calendario"])
    for row in plans:
        ws.append(row)
    ws = wb.create_sheet("Checklist")
    ws.append(["Código del plan", "Tarea"])
    for row in items:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("p.xlsx", buffer.getvalue())


def test_export_only_has_own_plans(client, tenant_a, tenant_b):
    with tenant_context(tenant_b.org):
        tenant_b.plan.name = "Plan secreto de B"
        tenant_b.plan.save()
    login(client, tenant_a.admin)
    wb = load_workbook(io.BytesIO(client.get(reverse("plan_export")).content))
    names = [row[1] for row in wb["Planes"].iter_rows(min_row=2, values_only=True)]
    assert names == ["Mensual"]


def test_import_cannot_reference_another_org_class(client, tenant_a, tenant_b):
    with tenant_context(tenant_b.org):
        tenant_b.asset.asset_class.code = "SOLO-B"
        tenant_b.asset.asset_class.save()
    login(client, tenant_a.admin)
    response = client.post(
        reverse("plan_import"),
        {"file": _upload([["N-1", "Nuevo", "SOLO-B", 1, "meses"]], [["N-1", "Revisar"]])},
    )
    row = response.context["preview"].rows[0]
    assert row.action == "error" and "no existe" in row.errors[0]


def test_same_code_in_b_is_a_new_plan_for_a_and_b_is_untouched(client, tenant_a, tenant_b):
    with tenant_context(tenant_b.org):
        tenant_b.plan.code = "SOLO-B"
        tenant_b.plan.save()
    login(client, tenant_a.admin)
    response = client.post(
        reverse("plan_import"),
        {"file": _upload([["SOLO-B", "De A", "", 1, "meses"]], [["SOLO-B", "Revisar"]])},
    )
    assert response.context["preview"].rows[0].action == "create"
    client.post(reverse("plan_import_confirm", args=[response.context["token"]]))
    with tenant_context(tenant_b.org):
        assert MaintenancePlan.objects.get(code="SOLO-B").name == "Mensual"
        assert PlanVersion.objects.count() == 1
    with tenant_context(tenant_a.org):
        assert MaintenancePlan.objects.get(code="SOLO-B").name == "De A"


def test_preview_token_of_a_is_rejected_for_b(client, tenant_a, tenant_b):
    login(client, tenant_a.admin)
    response = client.post(
        reverse("plan_import"), {"file": _upload([["N-2", "Nuevo", "", 1, "meses"]], [["N-2", "Revisar"]])}
    )
    token = response.context["token"]
    session = client.session
    pending = session[f"plan_import:{token}"]
    pending["organization"] = str(tenant_b.org.id)
    session[f"plan_import:{token}"] = pending
    session.save()
    client.post(reverse("plan_import_confirm", args=[token]))
    for tenant in (tenant_a, tenant_b):
        with tenant_context(tenant.org):
            assert not MaintenancePlan.objects.filter(code="N-2").exists()
