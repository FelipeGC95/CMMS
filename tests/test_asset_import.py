"""Importación de activos desde Excel (AST-08)."""

import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.accounts.models import MembershipLocation
from apps.assets.imports import (
    COLUMNS,
    SHEET_NAME,
    ImportFileError,
    apply_import,
    build_preview,
    build_template,
    parse_workbook,
)
from apps.assets.models import Asset, AssetClass, Meter
from apps.core.tenancy import tenant_context
from apps.maintenance.models import PlanAssignment
from apps.organizations.models import Location
from tests.conftest import login

HEADERS = [label for _, label, *_ in COLUMNS]


def xlsx(rows, headers=HEADERS, sheet=SHEET_NAME):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(
        "activos.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def row(code, name="Equipo", cls="Compresor", location="", parent="", crit="", meter="", reading="", **extra):
    values = {
        "code": code, "name": name, "asset_class": cls, "location": location, "parent": parent,
        "criticality": crit, "manufacturer": extra.get("maker", ""), "model": "", "serial_number": "",
        "year": extra.get("year", ""), "meter_kind": meter, "meter_value": reading,
    }  # fmt: skip
    return [values[key] for key, *_ in COLUMNS]


def preview_of(rows, membership=None):
    return build_preview(parse_workbook(xlsx(rows)), membership=membership)


# --- Lectura del archivo ----------------------------------------------------


def test_template_has_headers_dropdowns_and_reference_sheets(in_a):
    wb = load_workbook(io.BytesIO(build_template()))
    assert wb.sheetnames == [SHEET_NAME, "Instrucciones", "Clases y ubicaciones"]
    assert [c.value for c in wb[SHEET_NAME][1]] == HEADERS
    assert wb[SHEET_NAME].data_validations.dataValidation
    assert wb["Clases y ubicaciones"]["A2"].value == "Compresor"


def test_template_roundtrip_parses_the_example_row(in_a):
    upload = SimpleUploadedFile("t.xlsx", build_template())
    (example,) = parse_workbook(upload)
    assert example["code"] == "GRH-01" and example["year"] == "2019"


def test_missing_required_columns_is_a_file_error(in_a):
    with pytest.raises(ImportFileError, match="Clase"):
        parse_workbook(xlsx([["X", "Y"]], headers=["Código", "Nombre"]))


def test_headers_are_matched_without_accents_or_case(in_a):
    upload = xlsx([["A-1", "Bomba", "Bombas"]], headers=["CODIGO", "nombre", "Clase de activo"])
    (parsed,) = parse_workbook(upload)
    assert parsed == {"code": "A-1", "name": "Bomba", "asset_class": "Bombas", "_row": 2}


def test_non_excel_file_is_rejected(in_a):
    with pytest.raises(ImportFileError):
        parse_workbook(SimpleUploadedFile("x.xlsx", b"no soy un excel"))


# --- Vista previa -------------------------------------------------------------


def test_preview_classifies_rows_and_does_not_write(in_a):
    preview = preview_of(
        [
            row("CMP-01", "Compresor renombrado"),  # existe → actualiza
            row("NEW-1", cls="Bomba", location="Planta Norte"),  # nuevo con clase y sitio nuevos
            row("NEW-1"),  # repetido
            row("", "Sin código"),
            row("BAD-1", crit="Z", year="1800"),
            row("BAD-2", meter="km", reading="abc"),
            row("BAD-3", parent="NO-EXISTE"),
        ]
    )
    actions = [r.action for r in preview.rows]
    assert actions == ["update", "create", "error", "error", "error", "error", "error"]
    assert preview.new_classes == ["Bomba"] and preview.new_locations == ["Planta Norte"]
    assert "Criticidad inválida: usa A, B o C." in preview.rows[4].errors
    assert Asset.objects.count() == 1 and not AssetClass.objects.filter(name="Bomba").exists()


def test_new_classes_from_error_rows_are_not_created(in_a):
    preview = preview_of([row("E-1", cls="Clase fantasma", crit="Z")])
    assert preview.new_classes == []


def test_lower_meter_reading_on_update_is_skipped_with_note(in_a):
    meter = Meter.objects.create(asset=in_a.asset, kind="hours")
    from apps.assets.services import record_meter_reading

    record_meter_reading(meter=meter, value=500)
    preview = preview_of([row("CMP-01", meter="Horómetro", reading="400")])
    assert preview.rows[0].action == "update"
    assert "se omitirá" in preview.rows[0].notes[0]


# --- Aplicación -------------------------------------------------------------


def test_apply_creates_updates_and_is_idempotent(in_a):
    rows = [
        row("CMP-01", "Compresor principal", maker="Atlas Copco"),
        row("GRH-01", "Grúa 1", cls="Grúa horquilla", location="Bodega", crit="alta", meter="Horómetro",
            reading="4200", year="2019"),
        row("GRH-01-MAS", "Mástil", cls="Grúa horquilla", parent="GRH-01"),
    ]  # fmt: skip
    result = apply_import(preview_of(rows).to_session())
    assert (result.created, result.updated, result.classes_created, result.locations_created) == (2, 1, 1, 1)
    grua = Asset.objects.get(code="GRH-01")
    assert grua.criticality == "A" and grua.year == 2019 and grua.location.name == "Bodega"
    assert grua.meters.get().current_value == Decimal("4200")
    assert Asset.objects.get(code="GRH-01-MAS").parent == grua
    compresor = Asset.objects.get(code="CMP-01")
    assert compresor.name == "Compresor principal" and compresor.criticality == "A"  # vacío no borra

    again = apply_import(preview_of(rows).to_session())
    assert (again.created, again.updated, again.classes_created) == (0, 3, 0)
    assert Asset.objects.count() == 3


def test_apply_assigns_class_plans_to_new_assets_only(in_a):
    result = apply_import(preview_of([row("CMP-02"), row("CMP-01")]).to_session())
    assert result.plans_assigned == 1
    assert PlanAssignment.objects.filter(asset__code="CMP-02", plan=in_a.plan).exists()
    assert not PlanAssignment.objects.filter(asset__code="CMP-01").exists()

    result = apply_import(preview_of([row("CMP-03")]).to_session(), assign_class_plans=False)
    assert result.plans_assigned == 0


def test_generated_codes_for_new_classes_are_unique(in_a):
    apply_import(
        preview_of([row("A-1", cls="compresor de aire"), row("A-2", cls="Compresor-de aire!")]).to_session()
    )
    codes = list(AssetClass.objects.values_list("code", flat=True))
    assert len(codes) == len(set(c.lower() for c in codes))


# --- Alcance por sitio ------------------------------------------------------


def test_scoped_user_can_only_import_into_allowed_sites(in_a):
    allowed = Location.objects.create(code="PLT", name="Planta")
    child = Location.objects.create(code="L1", name="Línea 1", parent=allowed)
    other = Location.objects.create(code="OTR", name="Otra planta")
    planner = in_a.admin
    planner.restricted_to_locations = True
    planner.save()
    MembershipLocation.objects.create(membership=planner, location=allowed)

    preview = preview_of(
        [
            row("S-1", location="Planta"),
            row("S-2", location="L1"),  # descendiente: permitido
            row("S-3", location="Otra planta"),
            row("S-4"),  # sin ubicación: quedaría fuera de su vista
            row("S-5", location="Sitio inventado"),
            row("CMP-01", location="Planta"),  # existente sin ubicación permitida
        ],
        membership=planner,
    )
    assert [r.action for r in preview.rows] == ["create", "create", "error", "error", "error", "error"]
    assert preview.new_locations == []
    assert child.id and other.id


# --- Vistas -----------------------------------------------------------------


def test_import_flow_through_views(client, tenant_a):
    login(client, tenant_a.admin)
    assert client.get(reverse("asset_list")).content.decode().count(reverse("asset_import")) == 1
    template = client.get(reverse("asset_import_template"))
    assert template.status_code == 200 and template["Content-Type"].startswith("application/vnd.openxml")

    response = client.post(reverse("asset_import"), {"file": xlsx([row("NEW-9", "Bomba 9", cls="Bomba")])})
    assert response.status_code == 200 and response.context["counts"] == {
        "create": 1,
        "update": 0,
        "error": 0,
    }
    with tenant_context(tenant_a.org):
        assert not Asset.objects.filter(code="NEW-9").exists()

    token = response.context["token"]
    confirm = client.post(reverse("asset_import_confirm", args=[token]), {"assign_plans": "on"})
    assert confirm.status_code == 302 and confirm["Location"] == reverse("asset_list")
    with tenant_context(tenant_a.org):
        assert Asset.objects.filter(code="NEW-9", asset_class__name="Bomba").exists()
    # El token se usa una sola vez.
    reused = client.post(reverse("asset_import_confirm", args=[token]))
    assert reused["Location"] == reverse("asset_import")


def test_invalid_file_shows_error_in_form(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.post(reverse("asset_import"), {"file": SimpleUploadedFile("x.xlsx", b"nope")})
    assert response.status_code == 200 and response.context["form"].errors["file"]


@pytest.mark.parametrize("role", ["tech", "viewer"])
def test_only_asset_managers_can_import(client, tenant_a, role):
    login(client, getattr(tenant_a, role))
    assert client.get(reverse("asset_import")).status_code == 403
    assert client.get(reverse("asset_import_template")).status_code == 403
    assert reverse("asset_import") not in client.get(reverse("asset_list")).content.decode()
