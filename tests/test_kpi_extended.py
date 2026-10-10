"""KPI-02: indicadores ampliados con resultados conocidos, calculados a mano.

Ventanas fijas desde T0 = 1-sep-2026 (UTC). Capacidad del único técnico: 50 h/semana
(así h/50 nunca cae en un empate de redondeo a un decimal).
"""

import datetime
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import Membership
from apps.assets.models import Asset, AssetClass
from apps.core.tenancy import tenant_context
from apps.organizations.models import Location
from apps.reporting.extended import backlog_trend, cost_report, failure_pareto, work_mix
from apps.reporting.selectors import KpiFilters
from apps.work.models import (
    FailureCode,
    FailureEvent,
    LaborEntry,
    WorkOrder,
    WorkOrderTransition,
)
from apps.work.services import create_work_order
from tests.conftest import login

T0 = timezone.make_aware(datetime.datetime(2026, 9, 1))


def _at(days=0, hours=0):
    return T0 + datetime.timedelta(days=days, hours=hours)


def _wo(asset=None, *, work_type="corrective", minutes=None, created=None, location=None, **fields):
    wo = create_work_order(
        work_type=work_type, title="OT", asset=asset, location=location, estimated_minutes=minutes
    )
    updates = dict(fields)
    if created is not None:
        updates["created_at"] = created
    if updates:
        WorkOrder.objects.filter(pk=wo.pk).update(**updates)
        wo.refresh_from_db()
    return wo


def _cancel(wo, at):
    WorkOrder.objects.filter(pk=wo.pk).update(status=WorkOrder.Status.CANCELLED)
    transition = WorkOrderTransition.objects.create(
        work_order=wo, from_status=wo.status, to_status="cancelled"
    )
    WorkOrderTransition.objects.filter(pk=transition.pk).update(created_at=at)


# --- CA1 · Backlog en semanas ---------------------------------------------------------


@pytest.fixture
def backlog_data(in_a):
    Membership.objects.filter(pk=in_a.tech.pk).update(weekly_capacity_hours=Decimal("50"))
    asset = in_a.asset
    _wo(asset, minutes=600, created=_at(1))  # A: 10 h, sigue abierta
    _wo(asset, minutes=1200, created=_at(8), status="closed", completed_at=_at(16), closed_at=_at(16))  # B
    _cancel(_wo(asset, minutes=1800, created=_at(2)), _at(20))  # C: 30 h, cancelada el día 20
    _wo(asset, created=_at(3))  # D: sin estimación, abierta
    _wo(asset, minutes=120, created=_at(-10), status="closed", completed_at=_at(5), closed_at=_at(5))  # E
    _wo(asset, minutes=300, created=_at(22))  # F: 5 h, creada el día 22
    _wo(asset, minutes=999, created=_at(29))  # creada después de la ventana
    return in_a


def test_backlog_weeks_rebuilt_at_each_weekly_cut(backlog_data):
    trend = backlog_trend(membership=backlog_data.admin, filters=KpiFilters(days=28), now=_at(28))
    assert [p.at for p in trend.points] == [_at(7), _at(14), _at(21), _at(28)]
    assert [p.day for p in trend.points] == [
        datetime.date(2026, 9, 7),
        datetime.date(2026, 9, 14),
        datetime.date(2026, 9, 21),
        datetime.date(2026, 9, 28),
    ]
    assert [p.count for p in trend.points] == [3, 4, 2, 3]
    assert [p.estimated_hours for p in trend.points] == [
        Decimal("40.0"),
        Decimal("60.0"),
        Decimal("10.0"),
        Decimal("15.0"),
    ]
    assert [p.unestimated for p in trend.points] == [1, 1, 1, 1]
    assert trend.weekly_capacity_hours == Decimal("50") and trend.technicians == 1
    assert [p.weeks for p in trend.points] == [Decimal("0.8"), Decimal("1.2"), Decimal("0.2"), Decimal("0.3")]
    assert [p.bar_pct for p in trend.points] == [66.7, 100.0, 16.7, 25.0]


def test_backlog_without_capacity_shows_hours_not_weeks(backlog_data):
    Membership.objects.filter(pk=backlog_data.tech.pk).update(weekly_capacity_hours=0)
    trend = backlog_trend(membership=backlog_data.admin, filters=KpiFilters(days=28), now=_at(28))
    assert trend.by_hours and trend.weekly_capacity_hours == 0
    assert [p.weeks for p in trend.points] == [None, None, None, None]
    assert trend.current.estimated_hours == Decimal("15.0")


def test_backlog_series_defaults_to_twelve_weekly_points(in_a):
    trend = backlog_trend(membership=in_a.admin, filters=KpiFilters(days=84), now=_at(84))
    assert len(trend.points) == 12 and trend.points[-1].at == _at(84)
    assert all(p.count == 0 and p.weeks == Decimal("0.0") for p in trend.points)


# --- CA2 · Planificado vs. correctivo ----------------------------------------------------


def _done(asset, work_type, day, *, minutes=(), status="closed"):
    wo = _wo(asset, work_type=work_type, status=status, completed_at=_at(day))
    for m in minutes:
        LaborEntry.objects.create(work_order=wo, user=_user(), minutes=m)
    return wo


def _user():
    return Membership.objects.filter(role="technician").select_related("user").first().user


@pytest.fixture
def mix_data(in_a):
    asset = in_a.asset
    _done(asset, "preventive", 2, minutes=[120])  # sep: planificada 2 h
    _done(asset, "inspection", 9, minutes=[60])  # sep: planificada 1 h
    _done(asset, "predictive", 19, status="completed")  # sep: planificada sin horas
    _done(asset, "corrective", 4, minutes=[90, 30])  # sep: correctiva 2 h (dos registros)
    _done(asset, "improvement", 5, minutes=[45])  # sep: mejora, fuera del %
    _done(asset, "corrective", 6, minutes=[600], status="cancelled")  # cancelada: no cuenta
    _done(asset, "corrective", -3, minutes=[600])  # antes del período
    _wo(asset, work_type="corrective")  # abierta
    _done(asset, "preventive", 32, minutes=[60])  # oct: planificada 1 h
    return in_a


def test_work_mix_by_month_and_type(mix_data):
    filters = KpiFilters(date_from=datetime.date(2026, 9, 1), date_to=datetime.date(2026, 10, 10))
    mix = work_mix(membership=mix_data.admin, filters=filters, now=_at(60))
    sep, oct_ = mix.months
    assert (sep.month, oct_.month) == (datetime.date(2026, 9, 1), datetime.date(2026, 10, 1))
    assert (sep.planned_count, sep.corrective_count, sep.other_count) == (3, 1, 1)
    assert (sep.planned_minutes, sep.corrective_minutes, sep.other_minutes) == (180, 120, 45)
    assert sep.planned_pct == 60.0 and sep.planned_pct_count == 75.0
    assert (sep.bar_planned, sep.bar_corrective) == (60.0, 40.0)
    assert (oct_.planned_count, oct_.corrective_count, oct_.planned_pct) == (1, 0, 100.0)
    total = mix.total
    assert (total.planned_count, total.corrective_count, total.other_count) == (4, 1, 1)
    assert round(total.planned_pct, 4) == round(100 * 240 / 360, 4)
    assert total.planned_pct_count == 80.0
    assert total.other_hours == Decimal("0.8")


def test_work_mix_without_completed_orders_is_empty(in_a):
    mix = work_mix(membership=in_a.admin, filters=KpiFilters(days=28), now=_at(28))
    assert mix.total.planned_pct is None and mix.total.planned_pct_count is None
    assert all(not m.has_data and m.bar_planned == 0 for m in mix.months)


# --- CA3 · Costos por activo y por clase -------------------------------------------------


def _costs(asset, day, labor=0, parts=0, other=0, status="closed", location=None):
    total = Decimal(labor) + Decimal(parts) + Decimal(other)
    return _wo(
        asset,
        location=location,
        status=status,
        completed_at=_at(day),
        closed_at=_at(day) if status == "closed" else None,
        labor_cost=labor,
        parts_cost=parts,
        other_cost=other,
        total_cost=total,
    )


@pytest.fixture
def cost_data(in_a):
    compressor = in_a.asset.asset_class
    pumps = AssetClass.objects.create(code="BOM", name="Bomba")
    n1 = Asset.objects.create(code="N1", name="Compresor norte", asset_class=compressor)
    s1 = Asset.objects.create(code="S1", name="Bomba sur", asset_class=pumps)
    site = Location.objects.create(code="L", name="Bodega")
    _costs(in_a.asset, 2, labor=100, parts=50)
    _costs(in_a.asset, 4, labor=20, other=30)
    _costs(s1, 5, parts=300)
    _costs(n1, 6, labor=999, status="completed")  # completada, sin cerrar: totales no fijos
    _costs(n1, -1, labor=500)  # cerrada antes del período
    _costs(None, 7, other=40, location=site)  # OT sobre ubicación
    _costs(n1, 8)  # cerrada sin costos
    _costs(n1, 9, labor=70, status="cancelled")
    return {"tenant": in_a, "n1": n1, "s1": s1}


def test_costs_by_asset_and_class(cost_data):
    report = cost_report(membership=cost_data["tenant"].admin, filters=KpiFilters(days=28), now=_at(28))
    rows = [
        (r["asset__code"], r["work_orders"], r["labor"], r["parts"], r["other"], r["total"])
        for r in report.by_asset
    ]
    assert rows == [
        ("S1", 1, 0, 300, 0, 300),
        ("CMP-01", 2, 120, 50, 30, 200),
        ("N1", 1, 0, 0, 0, 0),
    ]
    classes = [
        (r["asset__asset_class__name"], r["assets"], r["work_orders"], r["total"], r["bar_pct"])
        for r in report.by_class
    ]
    assert classes == [("Bomba", 1, 1, 300, 100.0), ("Compresor", 2, 3, 200, 66.7)]
    assert report.location_totals == {
        "labor": 0,
        "parts": 0,
        "other": 40,
        "total": 40,
        "work_orders": 1,
    }
    assert report.totals == {"labor": 120, "parts": 350, "other": 70, "total": 540, "work_orders": 5}


def test_costs_follow_class_filter(cost_data):
    pumps = cost_data["s1"].asset_class
    report = cost_report(
        membership=cost_data["tenant"].admin, filters=KpiFilters(days=28, asset_class=pumps), now=_at(28)
    )
    assert [r["asset__code"] for r in report.by_asset] == ["S1"]
    assert report.totals["total"] == 300 and report.location_totals["work_orders"] == 0


def test_costs_without_closed_orders_are_zero(in_a):
    report = cost_report(membership=in_a.admin, filters=KpiFilters(days=28), now=_at(28))
    assert list(report.by_asset) == [] and report.by_class == []
    assert report.totals["total"] == 0 and report.totals["work_orders"] == 0


# --- CA4 · Pareto de fallas --------------------------------------------------------------


def _failure(asset, start, hours, *, impact="total_stop", mode=None, cause=None):
    return FailureEvent.objects.create(
        asset=asset,
        impact=impact,
        failure_mode=mode,
        failure_cause=cause,
        started_at=start,
        ended_at=start + datetime.timedelta(hours=hours) if hours is not None else None,
    )


@pytest.fixture
def failure_data(in_a):
    codes = {c: FailureCode.objects.create(kind="mode", code=c, name=f"Modo {c}") for c in ("ZA", "ZB", "ZC")}
    cause = FailureCode.objects.create(kind="cause", code="QA", name="Causa QA")
    asset = in_a.asset
    _failure(asset, _at(1), 2, mode=codes["ZA"], cause=cause)
    _failure(asset, _at(3), 4, mode=codes["ZA"], cause=cause)
    _failure(asset, _at(5), 5, impact="partial", mode=codes["ZA"], cause=cause)  # sin detención total
    _failure(asset, _at(6), 10, mode=codes["ZB"])
    _failure(asset, _at(27), 72, mode=codes["ZB"])  # termina después: se corta al fin (24 h)
    _failure(asset, _at(8), 3, impact="none", mode=codes["ZC"])
    _failure(asset, _at(9), 1)  # sin código
    _failure(asset, _at(27), None)  # sin código, sigue abierta: 24 h hasta el fin
    _failure(asset, _at(-2), 5, mode=codes["ZA"])  # antes del período
    return in_a


def test_pareto_by_count(failure_data):
    pareto = failure_pareto(membership=failure_data.admin, filters=KpiFilters(days=28), now=_at(28))
    rows = [(r.code, r.count, r.downtime_hours, round(r.cumulative_pct, 2), r.vital) for r in pareto.rows]
    assert rows == [
        ("ZA", 3, 6.0, 50.0, True),
        ("ZB", 2, 34.0, 83.33, True),  # cruza el 80 %: también cuenta
        ("ZC", 1, 0.0, 100.0, False),
    ]
    assert (pareto.total_count, pareto.vital_count) == (6, 2)
    assert (pareto.uncoded_count, pareto.uncoded_downtime_hours) == (2, 25.0)
    assert (pareto.rows[1].before_pct, pareto.rows[1].share_bar) == (50.0, 33.3)


def test_pareto_by_downtime(failure_data):
    pareto = failure_pareto(
        membership=failure_data.admin, filters=KpiFilters(days=28), metric="downtime", now=_at(28)
    )
    rows = [(r.code, r.value, r.cumulative_pct, r.vital) for r in pareto.rows]
    assert rows == [("ZB", 34.0, 85.0, True), ("ZA", 6.0, 100.0, False), ("ZC", 0.0, 100.0, False)]
    assert pareto.total_downtime_hours == 40.0


def test_pareto_by_cause_leaves_uncoded_apart(failure_data):
    pareto = failure_pareto(
        membership=failure_data.admin, filters=KpiFilters(days=28), code_kind="cause", now=_at(28)
    )
    assert [(r.code, r.count, r.cumulative_pct) for r in pareto.rows] == [("QA", 3, 100.0)]
    assert pareto.uncoded_count == 5


def test_pareto_without_failures(in_a):
    pareto = failure_pareto(membership=in_a.admin, filters=KpiFilters(days=28), now=_at(28))
    assert pareto.rows == [] and pareto.uncoded_count == 0 and pareto.vital_count == 0


# --- Pantalla y Excel (CA5, CA6) -----------------------------------------------------


def test_report_page_shows_formulas_and_accessible_tables(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.get(reverse("kpi_report"))
    assert response.status_code == 200
    body = response.content.decode()
    assert body.count('class="kpi__formula') == 4
    for reference in (
        "SMRP Best Practices (Backlog",
        "Planned Maintenance Percentage",
        "EN 15341",
        "ISO 14224",
    ):
        assert reference in body
    assert "style=" not in body and "<script>" not in body
    # Las cifras de los gráficos van sin localizar (atributos SVG con punto decimal).
    assert 'width="0,0%"' not in body


def test_report_default_period_is_twelve_weeks(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.get(reverse("kpi_report"))
    assert len(response.context["backlog"].points) == 12


def test_report_lists_top_assets_and_all_with_pagination(client, tenant_a):
    with tenant_context(tenant_a.org):
        for i in range(12):
            asset = Asset.objects.create(
                code=f"T{i:02d}", name=f"Equipo {i}", asset_class=tenant_a.asset.asset_class
            )
            _costs(asset, 0, labor=100 + i)
            WorkOrder.objects.filter(asset=asset).update(
                closed_at=timezone.now() - datetime.timedelta(days=1)
            )
    login(client, tenant_a.admin)
    top = client.get(reverse("kpi_report"))
    assert [r["asset__code"] for r in top.context["cost_rows"]][:2] == ["T11", "T10"]
    assert len(top.context["cost_rows"]) == 10 and top.context["more_assets"]
    every = client.get(reverse("kpi_report"), {"costs": "all"})
    assert len(every.context["cost_rows"]) == 12 and not every.context["more_assets"]


def test_excel_export_has_one_sheet_per_indicator_with_filters(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.get(reverse("kpi_report_export_xlsx"), {"period": "90", "pareto_by": "downtime"})
    assert response.status_code == 200
    assert "attachment;" in response["Content-Disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [
        "Filtros",
        "Backlog en semanas",
        "Preventivo-correctivo",
        "Costos por activo",
        "Costos por clase",
        "Pareto de fallas",
    ]
    filters = [list(r) for r in book["Filtros"].iter_rows(values_only=True)]
    assert filters[1][0] == "Período"
    assert len(list(book["Backlog en semanas"].iter_rows())) == 1 + 13  # 90 días: 13 cortes semanales
    assert next(book["Pareto de fallas"].iter_rows(values_only=True))[4] == "% de la detención"
