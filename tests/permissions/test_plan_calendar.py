"""Calendario de preventivos (PM-06): quién lo ve y alcance por sitio (ORG-06)."""

import io

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import Membership, MembershipLocation, User
from apps.assets.models import Asset
from apps.core.tenancy import tenant_context
from apps.maintenance.services import assign_plan
from apps.organizations.models import Location
from tests.conftest import PASSWORD, login

URLS = ["plan_calendar", "plan_calendar_export_xlsx"]


def _member(tenant, role):
    user = User.objects.create_user(f"{role}-cal@{tenant.org.slug}.cl", PASSWORD)
    return Membership.objects.create(user=user, organization=tenant.org, role=role)


@pytest.mark.parametrize("url_name", URLS)
@pytest.mark.parametrize("role", ["admin", "planner", "supervisor", "technician"])
def test_roles_that_work_orders_can_see_the_calendar(client, tenant_a, role, url_name):
    login(client, _member(tenant_a, role))
    assert client.get(reverse(url_name)).status_code == 200


@pytest.mark.parametrize("url_name", URLS)
@pytest.mark.parametrize("role", ["viewer", "external"])
def test_viewer_and_external_cannot_see_the_calendar(client, tenant_a, role, url_name):
    login(client, _member(tenant_a, role))
    assert client.get(reverse(url_name)).status_code == 403


@pytest.fixture
def scoped(tenant_a):
    """Técnico limitado a Planta Norte; un activo en su área y otro en Planta Sur, ambos con el plan."""
    with tenant_context(tenant_a.org):
        north = Location.objects.create(code="N", name="Planta Norte")
        area = Location.objects.create(code="N-A", name="Envasado", kind="area", parent=north)
        south = Location.objects.create(code="S", name="Planta Sur")
        cls = tenant_a.asset.asset_class
        inside = Asset.objects.create(code="IN-01", name="Llenadora", asset_class=cls, location=area)
        outside = Asset.objects.create(code="OUT-01", name="Caldera", asset_class=cls, location=south)
        for asset in (inside, outside):
            assign_plan(plan=tenant_a.plan, asset=asset, starts_on=timezone.localdate())
        MembershipLocation.objects.create(membership=tenant_a.tech, location=north)
    tenant_a.tech.restricted_to_locations = True
    tenant_a.tech.save()
    return north, inside, outside


def test_scoped_technician_sees_only_assets_of_their_sites(client, tenant_a, scoped):
    _north, inside, _outside = scoped
    login(client, tenant_a.tech)
    response = client.get(reverse("plan_calendar"), {"view": "year"})
    assert {o.asset for o in response.context["result"].occurrences} == {inside}
    export = client.get(reverse("plan_calendar_export_xlsx"))
    rows = list(load_workbook(io.BytesIO(export.content)).active.iter_rows(values_only=True))[1:]
    assert rows and {row[2] for row in rows} == {"IN-01"}


def test_location_filter_includes_descendants(client, tenant_a, scoped):
    north, inside, _outside = scoped
    login(client, tenant_a.admin)
    response = client.get(reverse("plan_calendar"), {"view": "year", "location": north.pk})
    assert {o.asset for o in response.context["result"].occurrences} == {inside}
