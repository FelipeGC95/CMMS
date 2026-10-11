"""Importación de repuestos y stock por Excel/CSV (INV-02)."""

import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.core.tenancy import tenant_context
from apps.inventory import services
from apps.inventory.imports import (
    COLUMNS,
    SHEET_NAME,
    ImportFileError,
    apply_import,
    build_preview,
    build_template,
    parse_file,
)
from apps.inventory.models import Part, StockItem, StockMovement
from tests.conftest import login

pytestmark = pytest.mark.usefixtures("reconciled_inventory")
HEADERS = [str(label) for _key, label, *_rest in COLUMNS]
KEYS = [key for key, *_rest in COLUMNS]


def row(code, name="Rodamiento", **values):
    data = {"code": code, "name": name, **values}
    return [data.get(key, "") for key in KEYS]


def xlsx(rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = str(SHEET_NAME)
    ws.append(headers)
    for values in rows:
        ws.append(values)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("repuestos.xlsx", buffer.getvalue())


def csv_file(text, encoding="utf-8"):
    return SimpleUploadedFile("repuestos.csv", text.encode(encoding), content_type="text/csv")


def run(membership, rows):
    preview = build_preview(parse_file(xlsx(rows)), membership=membership)
    return preview, (
        None if preview.has_errors else apply_import(preview.to_session(), membership=membership)
    )


# --- Archivo y plantilla --------------------------------------------------------


def test_template_roundtrip_parses_the_example_row(in_a):
    content = build_template(membership=in_a.admin)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames[0] == str(SHEET_NAME) and "Bodegas" in wb.sheetnames
    rows = parse_file(SimpleUploadedFile("plantilla.xlsx", content))
    assert rows[0]["code"] == "ROD-6205" and rows[0]["quantity"] == "12"


def test_csv_with_semicolons_and_latin1_is_accepted(in_a):
    rows = parse_file(csv_file("Código;Nombre;Cantidad\nFIL-1;Filtro aceite;3\n", encoding="cp1252"))
    assert rows == [{"code": "FIL-1", "name": "Filtro aceite", "quantity": "3", "_row": 2}]


def test_missing_required_columns_and_unreadable_files(in_a):
    with pytest.raises(ImportFileError, match="Faltan columnas obligatorias: Nombre"):
        parse_file(csv_file("codigo,cantidad\nA,1\n"))
    with pytest.raises(ImportFileError, match="No se pudo leer"):
        parse_file(SimpleUploadedFile("x.xlsx", b"no es excel"))
    with pytest.raises(ImportFileError, match="no tiene filas"):
        parse_file(csv_file("codigo,nombre\n"))


# --- Vista previa ---------------------------------------------------------------


def test_preview_reports_errors_per_row_and_writes_nothing(in_a):
    services.create_part(membership=in_a.admin, code="ROD-1", name="Rodamiento")
    preview = build_preview(
        parse_file(
            xlsx(
                [
                    row("rod-1", "", quantity=5),  # actualiza (sin nombre vale)
                    row("NEW-1", "Nuevo", quantity="-2"),
                    row("", "Sin código"),
                    row("NEW-2", "", min_quantity=5, target_quantity=2),
                    row("NEW-3", "X", warehouse="NO-EXISTE"),
                    row("NEW-4", "Y", quantity="1,5"),
                    row("NEW-4", "Y"),  # repetido en la misma bodega
                ]
            )
        ),
        membership=in_a.admin,
    )
    actions = [r.action for r in preview.rows]
    assert actions == ["update", "error", "error", "error", "error", "create", "error"]
    errors = " ".join(e for r in preview.rows for e in r.errors)
    for text in (
        "número mayor o igual a cero",
        "Falta el código",
        "Falta el nombre",
        "objetivo",
        "NO-EXISTE",
    ):
        assert text in errors
    assert "Conteo: 0 → 5." in preview.rows[0].notes
    assert preview.has_errors
    assert StockMovement.objects.count() == 0 and Part.objects.count() == 1


def test_apply_is_all_or_nothing(in_a):
    preview = build_preview(parse_file(xlsx([row("A", "A"), row("", "x")])), membership=in_a.admin)
    with pytest.raises(Exception, match="filas con errores"):
        apply_import(preview.to_session(), membership=in_a.admin)
    assert not Part.objects.exists()


def test_apply_rolls_back_everything_when_a_service_fails(in_a):
    preview = build_preview(
        parse_file(xlsx([row("A", "A", quantity=2), row("B", "B", quantity=3)])), membership=in_a.admin
    )
    session = preview.to_session()
    session["rows"][1]["warehouse_id"] = "00000000-0000-7000-8000-000000000000"  # bodega que no existe
    with pytest.raises(Exception, match="Bodega no encontrada"):
        apply_import(session, membership=in_a.admin)
    assert not Part.objects.exists() and not StockMovement.objects.exists()


# --- Aplicación: conteo, costo e idempotencia ----------------------------------


def test_apply_creates_parts_counts_stock_and_sets_levels(in_a):
    _preview, result = run(
        in_a.admin,
        [
            row(
                "ROD-1",
                "Rodamiento 6205",
                unit="un",
                category="Rodamientos",
                manufacturer_ref="6205",
                quantity=10,
                unit_cost=1500,
                min_quantity=4,
                target_quantity=12,
                bin_location="A-3",
            ),
            row("FIL-1", "Filtro", quantity=0),
        ],
    )
    assert (result.created, result.updated, result.adjusted, result.unchanged) == (2, 0, 1, 1)
    part = Part.objects.get(code="ROD-1")
    assert part.category == "Rodamientos" and part.avg_cost == Decimal(1500)
    item = StockItem.objects.get(part=part)
    assert (item.on_hand, item.min_quantity, item.target_quantity, item.bin_location) == (
        Decimal(10),
        Decimal(4),
        Decimal(12),
        "A-3",
    )
    movement = StockMovement.objects.get(part=part)
    assert movement.kind == "adjustment" and movement.reason == "import"


def test_reimporting_the_same_file_is_idempotent(in_a):
    rows = [row("ROD-1", "Rodamiento", quantity=10, unit_cost=1500, min_quantity=4)]
    run(in_a.admin, rows)
    preview, result = run(in_a.admin, rows)
    assert [r.action for r in preview.rows] == ["update"]
    assert "Stock sin cambios (10)." in preview.rows[0].notes
    assert (result.created, result.adjusted, result.unchanged) == (0, 0, 1)
    assert StockMovement.objects.count() == 1


def test_quantity_is_a_count_and_cost_only_sets_average_without_previous_stock(in_a):
    run(in_a.admin, [row("ROD-1", "Rodamiento", quantity=10, unit_cost=100)])
    run(in_a.admin, [row("ROD-1", "", quantity=4, unit_cost=999)])  # baja: no cambia el promedio
    part = Part.objects.get(code="ROD-1")
    assert StockItem.objects.get(part=part).on_hand == 4
    assert part.avg_cost == Decimal(100)
    assert [m.quantity for m in StockMovement.objects.order_by("created_at")] == [Decimal(10), Decimal(-6)]


def test_blank_cells_do_not_erase_data_and_levels_keep_previous_values(in_a):
    run(in_a.admin, [row("ROD-1", "Rodamiento", category="Rodamientos", min_quantity=4, bin_location="A-3")])
    run(in_a.admin, [row("ROD-1", "", target_quantity=9)])
    part = Part.objects.get(code="ROD-1")
    item = StockItem.objects.get(part=part)
    assert part.name == "Rodamiento" and part.category == "Rodamientos"
    assert (item.min_quantity, item.target_quantity, item.bin_location) == (Decimal(4), Decimal(9), "A-3")


def test_rows_for_two_warehouses_of_the_same_part(in_a):
    services.ensure_default_warehouse()
    services.create_warehouse(membership=in_a.admin, code="NOR", name="Bodega norte")
    preview, result = run(
        in_a.admin, [row("ROD-1", "Rodamiento", quantity=3), row("ROD-1", "", warehouse="nor", quantity=2)]
    )
    assert preview.count("create") == 1
    assert (result.created, result.adjusted) == (1, 2)
    assert sorted(StockItem.objects.values_list("warehouse__code", "on_hand")) == [
        ("BOD", Decimal(3)),
        ("NOR", Decimal(2)),
    ]


# --- Vistas ---------------------------------------------------------------------


def test_import_flow_through_views(client, tenant_a):
    login(client, tenant_a.admin)
    assert client.get(reverse("inventory_import")).status_code == 200
    template = client.get(reverse("inventory_import_template"))
    assert template["Content-Type"].startswith("application/vnd.openxmlformats")
    response = client.post(
        reverse("inventory_import"), {"file": xlsx([row("ROD-1", "Rodamiento", quantity=5)])}
    )
    assert response.status_code == 200 and "Revisa antes de guardar" in response.content.decode()
    token = response.context["token"]
    response = client.post(reverse("inventory_import_confirm", args=[token]), follow=True)
    assert "Importación lista: 1 repuesto creado" in response.content.decode()
    with tenant_context(tenant_a.org):
        assert StockItem.objects.get(part__code="ROD-1").on_hand == 5
    # La vista previa se usa una sola vez.
    response = client.post(reverse("inventory_import_confirm", args=[token]), follow=True)
    assert "La vista previa expiró" in response.content.decode()


def test_preview_with_errors_cannot_be_confirmed(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.post(reverse("inventory_import"), {"file": xlsx([row("A", "A"), row("", "B")])})
    body = response.content.decode()
    assert response.context["token"] == ""
    assert "todo o nada" in body and "Importar 2 filas" not in body


def test_invalid_file_shows_error_in_form(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.post(reverse("inventory_import"), {"file": SimpleUploadedFile("x.xlsx", b"nada")})
    assert response.status_code == 200
    assert response.context["form"].errors["file"]


def test_preview_token_of_another_organization_is_rejected(client, tenant_a, tenant_b):
    login(client, tenant_b.admin)
    response = client.post(reverse("inventory_import"), {"file": xlsx([row("B-1", "De B", quantity=1)])})
    token = response.context["token"]
    session = client.session
    data = session[f"part_import:{token}"]
    data["organization"] = str(tenant_a.org.pk)
    session[f"part_import:{token}"] = data
    session.save()
    response = client.post(reverse("inventory_import_confirm", args=[token]), follow=True)
    assert "La vista previa expiró o es de otra organización" in response.content.decode()
    with tenant_context(tenant_b.org):
        assert not Part.objects.exists()
