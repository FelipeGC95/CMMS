"""Pantalla del calendario de preventivos (PM-06 CA1): mes, 12 meses y Excel."""

import io

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.tenancy import tenant_context
from apps.maintenance.models import add_months
from apps.maintenance.services import add_plan_absorption, assign_plan, create_plan
from tests.conftest import login

ITEM = [{"text": "Revisar", "item_type": "check"}]


def _setup(tenant):
    """Semestral que absorbe al mensual, en el activo de A, desde el día 3 del mes siguiente."""
    starts = add_months(timezone.localdate().replace(day=3), 1)
    with tenant_context(tenant.org):
        big = create_plan(
            code="SEM", name="Semestral", calendar_interval=6, calendar_unit="months",
            estimated_minutes=240, items=ITEM,
        )  # fmt: skip
        small = create_plan(
            code="MEN", name="Lubricación mensual", calendar_interval=1, calendar_unit="months",
            estimated_minutes=30, items=ITEM,
        )  # fmt: skip
        add_plan_absorption(membership=tenant.admin, plan=big, absorbed_plan=small)
        assign_plan(plan=big, asset=tenant.asset, starts_on=starts)
        assign_plan(plan=small, asset=tenant.asset, starts_on=starts)
    return starts, big, small


def test_month_view_shows_projection_and_absorption(client, tenant_a):
    starts, _big, _small = _setup(tenant_a)
    login(client, tenant_a.admin)
    response = client.get(reverse("plan_calendar"), {"month": f"{starts:%Y-%m}"})
    assert response.status_code == 200
    html = response.content.decode()
    assert "Semestral" in html and "Absorbida en Semestral" in html
    assert f'id="d-{starts:%Y-%m-%d}"' in html
    assert "month-grid" in html


def test_month_param_is_clamped_to_the_twelve_month_window(client, tenant_a):
    login(client, tenant_a.admin)
    first = timezone.localdate().replace(day=1)
    response = client.get(reverse("plan_calendar"), {"month": "1999-01"})
    assert response.context["month"] == first
    response = client.get(reverse("plan_calendar"), {"month": "2999-01"})
    assert response.context["month"] == add_months(first, 11)
    assert response.context["next_query"] == ""
    response = client.get(reverse("plan_calendar"), {"month": "nada"})
    assert response.context["month"] == first and response.context["prev_query"] == ""


def test_year_view_sums_counts_and_hours(client, tenant_a):
    _setup(tenant_a)
    login(client, tenant_a.admin)
    response = client.get(reverse("plan_calendar"), {"view": "year"})
    assert response.status_code == 200
    months = response.context["months"]
    assert len(months) == 12
    # Meses 2 a 12: 2 semestrales (con el mensual absorbido) + 9 mensuales sueltos.
    assert sum(m["count"] for m in months) == 11
    assert sum(m["absorbed"] for m in months) == 2
    assert sum(m["minutes"] for m in months) == 2 * 270 + 9 * 30
    assert response.context["total_hours"] == "13,5"


def test_filters_by_plan_and_class(client, tenant_a):
    _starts, _big, small = _setup(tenant_a)
    login(client, tenant_a.admin)
    response = client.get(reverse("plan_calendar"), {"view": "year", "plan": small.pk})
    assert sum(m["count"] for m in response.context["months"]) == 9
    assert sum(m["absorbed"] for m in response.context["months"]) == 2
    with tenant_context(tenant_a.org):
        from apps.assets.models import AssetClass

        other = AssetClass.objects.create(code="OTR", name="Otra")
    response = client.get(reverse("plan_calendar"), {"view": "year", "asset_class": other.pk})
    assert sum(m["count"] for m in response.context["months"]) == 0


def test_export_lists_the_twelve_month_projection(client, tenant_a):
    _setup(tenant_a)
    login(client, tenant_a.admin)
    response = client.get(reverse("plan_calendar_export_xlsx"))
    assert response.status_code == 200
    assert response["Content-Disposition"].endswith('.xlsx"')
    rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0][:3] == ("Fecha", "Plan", "Activo")
    assert len(rows) == 1 + 13
    assert any(r[1] == "Lubricación mensual" and r[8] == "Absorbida en Semestral" for r in rows[1:])
    assert {r[5] for r in rows[1:]} == {"Proyectada"}


def test_plan_list_and_menu_link_to_the_calendar(client, tenant_a):
    login(client, tenant_a.admin)
    url = reverse("plan_calendar")
    assert url in client.get(reverse("plan_list")).content.decode()
    assert url in client.get(reverse("menu")).content.decode()
    login(client, tenant_a.tech)
    assert url in client.get(reverse("menu")).content.decode()
    login(client, tenant_a.viewer)
    assert url not in client.get(reverse("menu")).content.decode()


def test_month_view_marks_started_work_orders(client, tenant_a):
    from apps.maintenance.services import generate_due_work_orders

    starts, _big, _small = _setup(tenant_a)
    with tenant_context(tenant_a.org):
        [wo] = generate_due_work_orders(today=starts, horizon_days=1)
        wo.status = "in_progress"
        wo.save(update_fields=["status"])
    login(client, tenant_a.admin)
    html = client.get(reverse("plan_calendar"), {"month": f"{starts:%Y-%m}"}).content.decode()
    assert "month-grid__item--started" in html
    assert "En ejecución" in html and wo.code in html
