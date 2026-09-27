"""Planes por Excel (PM-08): exportar la versión vigente e importar con vista previa y errores por fila."""

import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.core.tenancy import tenant_context
from apps.maintenance.models import MaintenancePlan, PlanVersion
from apps.maintenance.plan_excel import (
    apply_plan_import,
    build_plan_preview,
    build_plans_workbook,
    parse_plan_workbook,
)
from apps.maintenance.services import add_draft_item, create_draft
from tests.conftest import login

PLAN_HEADER = [
    "Código", "Nombre", "Clase", "Cada (calendario)", "Unidad de calendario", "Medidor", "Cada (medidor)",
    "Reprogramación", "Tolerancia (%)", "Duración estimada (min)", "Instrucciones",
]  # fmt: skip
ITEM_HEADER = ["Código del plan", "Orden", "Tarea", "Tipo", "Unidad", "Mínimo", "Máximo", "Obligatorio"]
RESOURCE_HEADER = ["Código del plan", "Tipo", "Descripción", "Cantidad", "Unidad"]


def _file(plans, items, resources=(), name="planes.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Planes"
    ws.append(PLAN_HEADER)
    for row in plans:
        ws.append(row)
    ws = wb.create_sheet("Checklist")
    ws.append(ITEM_HEADER)
    for row in items:
        ws.append(row)
    ws = wb.create_sheet("Recursos")
    ws.append(RESOURCE_HEADER)
    for row in resources:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def _preview(tenant, upload):
    return build_plan_preview(parse_plan_workbook(upload), membership=tenant.admin)


def test_export_has_current_version_items_and_resources(in_a):
    content = build_plans_workbook(membership=in_a.admin)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames[:3] == ["Planes", "Checklist", "Recursos"]
    plans = list(wb["Planes"].iter_rows(values_only=True))
    assert plans[1][:5] == ("PM1", "Mensual", "CMP", 30, "días")
    items = list(wb["Checklist"].iter_rows(values_only=True))
    assert items[2] == ("PM1", 2, "Presión", "Medición", "bar", 6.5, 8, "Sí")


def test_export_round_trip_is_unchanged(in_a):
    upload = SimpleUploadedFile("x.xlsx", build_plans_workbook(membership=in_a.admin))
    preview = _preview(in_a, upload)
    assert [(r.code, r.action) for r in preview.rows] == [("PM1", "unchanged")]


def test_new_code_creates_published_plan(in_a):
    upload = _file(
        [
            [
                "LUB-01",
                "Lubricación",
                "CMP",
                None,
                None,
                "Horómetro",
                250,
                "Fijo",
                5,
                45,
                "Con equipo detenido",
            ]
        ],
        [
            ["LUB-01", 1, "Engrasar rodamientos", "Verificación", None, None, None, "Sí"],
            ["LUB-01", 2, "Temperatura", "Medición", "°C", "40", "80,5", "No"],
            ["LUB-01", 3, "Foto del punto", "Foto", None, None, None, "No"],
        ],
        [["LUB-01", "Repuesto", "Grasa EP2", "0,5", "kg"]],
    )
    preview = _preview(in_a, upload)
    [row] = preview.rows
    assert row.action == "create" and row.frequency_label == "250 h" and row.item_count == 3
    result = apply_plan_import(preview.to_session(), membership=in_a.admin)
    assert result.created == 1
    plan = MaintenancePlan.objects.get(code="LUB-01")
    version = plan.current_version
    assert plan.asset_class.code == "CMP" and version.status == PlanVersion.Status.PUBLISHED
    assert (version.meter_kind, version.meter_interval, version.schedule_mode) == ("hours", 250, "fixed")
    assert (version.tolerance_pct, version.estimated_minutes) == (5, 45)
    temperature = version.items.get(order=2)
    assert (temperature.min_value, temperature.max_value, temperature.required) == (
        40,
        Decimal("80.5"),
        False,
    )
    assert version.items.get(order=3).required  # la foto siempre es obligatoria
    assert version.resources.get().quantity == Decimal("0.5")


def test_existing_code_with_changes_creates_a_draft_and_keeps_published(in_a):
    published = in_a.plan.current_version
    upload = _file(
        [["PM1", "Mensual", "CMP", 2, "meses", None, None, None, None, None, None]],
        [
            ["PM1", 1, "Purgar", None, None, None, None, None],
            ["PM1", 2, "Revisar filtro", None, None, None, None, None],
        ],
    )
    preview = _preview(in_a, upload)
    assert preview.rows[0].action == "draft"
    result = apply_plan_import(preview.to_session(), membership=in_a.admin)
    assert result.drafts == 1
    in_a.plan.refresh_from_db()
    assert in_a.plan.current_version == published  # la publicada no cambia
    assert published.items.count() == 2 and published.calendar_interval == 30
    draft = in_a.plan.versions.get(status=PlanVersion.Status.DRAFT)
    assert (draft.number, draft.calendar_interval, draft.calendar_unit) == (2, 2, "months")
    assert list(draft.items.values_list("text", flat=True)) == ["Purgar", "Revisar filtro"]


def test_row_errors_are_reported_and_skipped(in_a, tenant_b):
    create_draft(membership=in_a.admin, plan=in_a.plan)
    upload = _file(
        [
            ["PM1", "Mensual", None, 7, "días", None, None, None, None, None, None],  # ya tiene borrador
            ["X-1", "Sin frecuencia", None, None, None, None, None, None, None, None, None],
            ["X-2", "Clase ajena", "OTRA", 1, "meses", None, None, None, None, None, None],
            ["X-3", "Unidad mala", None, 1, "lunas", None, None, None, None, None, None],
            ["X-4", "Rango malo", None, 1, "semanas", None, None, None, None, None, None],
            ["X-5", "Sin checklist", None, 1, "semanas", None, None, None, None, None, None],
            ["X-4", "Repetido", None, 1, "semanas", None, None, None, None, None, None],
            ["OK-1", "Correcto", None, 1, "semanas", None, None, None, None, None, None],
        ],
        [
            ["PM1", 1, "Purgar", None, None, None, None, None],
            ["X-1", 1, "A", None, None, None, None, None],
            ["X-2", 1, "A", None, None, None, None, None],
            ["X-3", 1, "A", None, None, None, None, None],
            ["X-4", 1, "Presión", "Medición", "bar", 9, 1, None],
            ["OK-1", 1, "A", None, None, None, None, None],
            ["NADIE", 1, "A", None, None, None, None, None],
        ],
    )
    preview = _preview(in_a, upload)
    by_code = {}
    for row in preview.rows:
        by_code.setdefault(row.code, []).append(row)
    assert "borrador" in by_code["PM1"][0].errors[0]
    assert any("frecuencia" in e for e in by_code["X-1"][0].errors)
    assert any("no existe" in e for e in by_code["X-2"][0].errors)
    assert any("lunas" in e for e in by_code["X-3"][0].errors)
    assert any("mínimo" in e for e in by_code["X-4"][0].errors)
    assert any("se repite" in e for e in by_code["X-4"][1].errors)
    assert any("Checklist" in e for e in by_code["X-5"][0].errors)
    assert by_code["OK-1"][0].action == "create"
    assert preview.count("error") == 7 and len(preview.valid_rows) == 1
    assert "NADIE" in preview.orphans[0]
    apply_plan_import(preview.to_session(), membership=in_a.admin)
    assert set(MaintenancePlan.objects.values_list("code", flat=True)) == {"PM1", "OK-1"}


def test_web_flow_upload_preview_confirm(client, tenant_a):
    login(client, tenant_a.admin)
    assert client.get(reverse("plan_import")).status_code == 200
    export = client.get(reverse("plan_export"))
    assert export["Content-Type"].startswith("application/vnd.openxmlformats")
    upload = _file(
        [["NEW-1", "Nuevo", None, 1, "meses", None, None, None, None, None, None]],
        [["NEW-1", 1, "Revisar", None, None, None, None, None]],
    )
    response = client.post(reverse("plan_import"), {"file": upload})
    assert response.status_code == 200 and "NEW-1" in response.content.decode()
    token = response.context["token"]
    with tenant_context(tenant_a.org):
        assert not MaintenancePlan.objects.filter(code="NEW-1").exists()
    response = client.post(reverse("plan_import_confirm", args=[token]), follow=True)
    assert "1 plan creado" in response.content.decode()
    with tenant_context(tenant_a.org):
        assert MaintenancePlan.objects.get(code="NEW-1").current_version.number == 1
    # El token se usa una sola vez.
    response = client.post(reverse("plan_import_confirm", args=[token]))
    assert response.status_code == 302 and response.url == reverse("plan_import")


def test_bad_file_is_rejected_with_a_message(client, tenant_a):
    login(client, tenant_a.admin)
    response = client.post(reverse("plan_import"), {"file": SimpleUploadedFile("x.xlsx", b"no es excel")})
    assert "No se pudo leer el archivo" in response.content.decode()


def test_draft_with_existing_item_blocks_import(in_a):
    draft = create_draft(membership=in_a.admin, plan=in_a.plan)
    add_draft_item(membership=in_a.admin, version=draft, text="Nuevo", item_type="check")
    upload = SimpleUploadedFile("x.xlsx", build_plans_workbook(membership=in_a.admin))
    # Con un borrador abierto el plan no se importa: se informa para no pisar ese trabajo.
    row = _preview(in_a, upload).rows[0]
    assert row.action == "error" and "borrador" in row.errors[0]
